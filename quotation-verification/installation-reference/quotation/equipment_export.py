"""Customer exports contain actual supplied specifications, never internal costs."""
import io
import openpyxl
from openpyxl.styles import Alignment, Font, Border, Side
from openpyxl.utils import get_column_letter
from .installation import installation_enabled


def public_specification(row, result):
    point = result.get('detail', {}).get('selected_fan')
    if point:
        return f"{point['model']}；功率{'/'.join(point['powers'])}kW；风量{point['airflow']}m³/h；风压{point['pressure']}Pa；转速{point['speed']}r/min"
    text = row.get('spec', '')
    a = result.get('attributes', row.get('attributes', {}))
    if not text:
        if a.get('diameter'): text = 'Φ'+str(a['diameter'])
        else: text = '×'.join(str(a[k]) for k in ('width', 'height', 'length') if a.get(k))
        if a.get('model'): text = str(a['model']) + ('；'+text if text else '')
    return text


def public_notes(body, result):
    options = body.get('equipment_options', {})
    tax = '含税' if options.get('tax_included') else '不含税'
    freight = '含运费' if options.get('include_freight') else '不含运费'
    if body.get('pricing_mode') == 'mixed_v1':
        duct_freight = '含运费' if body.get('duct_options', {}).get('include_freight') else '不含运费'
        freight = f'设备{freight}；风管{duct_freight}'
    included = installation_enabled(body)
    notes = [f'报价范围：材料{tax}；{freight}；' + ('含安装（安装统一含税）。' if included else '不含安装。')]
    if result.get('payment'):
        notes.append('付款条件：'+result['payment']['label']+'；整单付款系数'+str(result['payment']['coefficient'])+'，材料与安装合计后统一乘一次并向上取整数元；明细单价为付款前价格。')
    notes += ['风机按明细所列实际供货型号、功率及工况报价。', '实际数量、规格及工艺，以双方确认的图纸、清单或订单为准。']
    if included:
        notes[-1] += ' 安装不含控制箱及相关电线电器、风机基础预制、抗震支架、预留洞开启与修复。'
    if body.get('note'): notes.append('报价备注：'+body['note'])
    return notes


def write_equipment_draft(body, result, identifier):
    from .excel_io import _formal_duct_lines, _excel_number
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = '报价草稿'
    records = [
        ['科庆 · 报价草稿（待核对）'],
        [f"客户：{body.get('customer', '')}；项目：{body.get('project', '')}；编号：{identifier}"],
        ['；'.join(result.get('issues', [])) or '项目已试算，请核对供货范围'],
        ['序号', '产品名称', '供货规格/型号', '单位', '数量', '单价（元）', '金额（元）', '状态/备注'],
    ]
    # Retain pending/excluded original rows; reuse formal line expansion for duct accessories.
    for i, (row, calculated) in enumerate(zip(body['rows'], result['rows']), 1):
        if calculated['status'] == 'priced':
            partial = {'rows': [calculated]}
            lines = _formal_duct_lines({**body, 'rows': [row]}, partial, False)
            for item in lines:
                records.append([str(i), item['name'], item['spec'], item['unit'], item['quantity'], item['material_unit'], item['material_amount'], item['remark']])
        else:
            records.append([str(i), row.get('name', ''), row.get('spec', ''), row.get('unit', ''), row.get('quantity', ''),
                            '待确认' if calculated['status'] == 'pending' else '', '', calculated.get('reason', '')])
    extras = _formal_duct_lines({**body, 'rows': []}, {**result, 'rows': []}, installation_enabled(body))
    for item in extras:
        records.append(['', item['name'], item['spec'], item['unit'], item['quantity'], item['material_unit'] if item['material_unit'] is not None else item['installation_unit'], item['material_amount'] if item['material_amount'] is not None else item['installation_amount'], item['remark']])
    for line in result.get('fees', {}).get('installation', {}).get('lines', []):
        if line['status'] == 'pending':
            records.append(['', line['name']+'安装', '', line.get('unit', ''), line.get('quantity', ''), '待确认', '', line['reason']])
    records += [['付款前材料小计', '', '', '', '', '', _excel_number(result['subtotal'])]]
    if installation_enabled(body):
        records += [['含税安装小计', '', '', '', '', '', _excel_number(result.get('fees', {}).get('installation', {}).get('amount'))]]
    records += [['试算总额', '', '', '', '', '', _excel_number(result['total']) if result.get('total') is not None else '待确认']]
    records += [[note] for note in public_notes(body, result)]
    for values in records:
        ws.append(values)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str): cell.data_type = 's'
            cell.font = Font(name='微软雅黑', size=10, bold=cell.row in {1, 4})
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            cell.border = Border(bottom=Side(style='hair', color='DDDDDD'))
            if cell.column in {6, 7} and isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0' if cell.value == int(cell.value) else '#,##0.########'
        ws.row_dimensions[ws.max_row].height = 42 if len(values) > 1 else 30
        if len(values) == 1: ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=8)
        elif values[0] in {'付款前材料小计', '含税安装小计', '试算总额'}:
            ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=6)
    for i, width in enumerate([8, 25, 65, 10, 12, 16, 18, 48], 1): ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = None
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    out = io.BytesIO(); wb.save(out); wb.close()
    return out.getvalue()
