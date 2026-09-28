import io

import openpyxl

from quotation.excel_io import (_formal_specification, build_quote_rows, suggest_duct_mapping,
                                extract_duct_flange_meter_totals, read_workbook, write_quote)
from quotation.store import dumps
from test_duct_pricing import catalog


def test_formal_specification_keeps_size_and_removes_repeated_thickness():
    assert _formal_specification({"spec": "0.75 / 800×320", "attributes": {}}) == "800×320"
    assert _formal_specification({"spec": "δ=0.75mm", "attributes": {"model": "800×320"}}) == "800×320"


def test_explicit_engineering_quantity_header_is_suggested():
    sheet = {"name": "清单", "rows": [
        {"number": 1, "cells": ["产品名称", "规格", "工程量", "数量"], "issues": []},
        {"number": 2, "cells": ["镀锌风管", "0.75mm", "100", "12"], "issues": []},
    ]}
    suggestion = suggest_duct_mapping(sheet)
    assert suggestion == {"mapping": {"name": 0, "spec": 1, "engineering_quantity": 2}, "start_row": 2}
    rows = build_quote_rows(sheet, 2, suggestion["mapping"], "duct_v1")
    assert rows[0]["engineering_quantity"] == "100"
    assert rows[0]["quantity"] == ""


def test_missing_or_ambiguous_area_column_stays_unpriced_for_manual_entry():
    sheet = {"name": "清单", "rows": [
        {"number": 1, "cells": ["产品名称", "工程量", "工程量"], "issues": []},
        {"number": 2, "cells": ["镀锌风管", "100", "101"], "issues": []},
    ]}
    suggestion = suggest_duct_mapping(sheet)
    assert "engineering_quantity" not in suggestion["mapping"]
    rows = build_quote_rows(sheet, 2, {"name": 0}, "duct_v1")
    assert rows[0]["engineering_quantity"] == ""


def test_split_duct_format_and_whole_quote_flange_rows_are_recognized():
    sheet = {"name": "Sheet2", "rows": [
        {"number": 1, "cells": ["序号", "构件名称", "规格范围", "法兰方式", "规格型号", "单位", "清单总量"], "issues": []},
        {"number": 2, "cells": ["2", "镀锌风管", "中压 δ=0.75", "角铁", "1800×630", "m2", "12.5"], "issues": []},
        {"number": 3, "cells": ["辅材", "30#角钢", "", "", "", "米", "9200"], "issues": []},
        {"number": 4, "cells": ["", "40#角钢", "", "", "", "米", "25000"], "issues": []},
    ]}
    suggestion = suggest_duct_mapping(sheet)
    assert suggestion == {"mapping": {"name": 1, "engineering_quantity": 6, "unit": 5, "flange": 3,
                                       "spec": 2, "model": 4}, "start_row": 2}
    flange_data = extract_duct_flange_meter_totals(sheet, 2, suggestion["mapping"])
    assert flange_data["totals"] == {"flange_30_meters": "9200", "flange_40_meters": "25000"}
    rows = build_quote_rows(sheet, 2, suggestion["mapping"], "duct_v1", flange_data["source_rows"])
    assert len(rows) == 1
    assert rows[0]["spec"] == "中压 δ=0.75"
    assert rows[0]["engineering_quantity"] == "12.5"
    assert rows[0]["attributes"] == {"model": "1800×630", "flange": "angle"}


def test_imported_mixed_flanges_are_shown_as_mixed_and_kept_per_row(admin):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["序号", "构件名称", "规格范围", "法兰方式", "规格型号", "单位", "清单总量"])
    sheet.append([1, "镀锌风管", "δ=0.6", "共板", "400×200", "m2", 10])
    sheet.append([2, "镀锌风管", "δ=1.2", "角铁", "1400×320", "m2", 20])
    sheet.append([3, "镀锌风管", "δ=0.75", "圆法兰", "Φ500", "m2", 30])
    sheet.append(["辅材", "30#角钢", "", "支", "", "米", 100])
    stream = io.BytesIO()
    workbook.save(stream)

    uploaded = admin.client.post(
        "/api/imports",
        data={"file": (io.BytesIO(stream.getvalue()), "混合法兰.xlsx")},
        headers={"X-CSRF-Token": admin.csrf},
        content_type="multipart/form-data",
    )
    assert uploaded.status_code == 200, uploaded.json
    created = admin.call("/quotes", "POST", {
        "pricing_mode": "duct_v1", "import_id": uploaded.json["id"], "sheet": "Sheet1", "start_row": 2,
        "mapping": {"name": 1, "spec": 2, "flange": 3, "model": 4, "unit": 5, "engineering_quantity": 6},
    })
    assert created.status_code == 200, created.json
    body = admin.call("/quotes/" + created.json["id"]).json["body"]
    assert body["duct_options"]["flange"] == "mixed"
    assert body["duct_options"]["flange_30_meters"] == "100"
    assert [row["attributes"]["flange"] for row in body["rows"]] == ["common", "angle", "round"]


def test_completed_quote_format_propagates_merged_state_and_extracts_accessories():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["序号", "构件名称", "规格范围", "法兰方式", "规格型号", "单位", "清单总量",
                  "单价", "金额", "报价备注", "备注"])
    sheet.append([1, "镀锌风管", "δ=0.75", "共板", "500×320", "m2", 12, "=1+1", "=G2*H2",
                  "40g半成品不含税", "一层"])
    sheet.append([2, "镀锌风管", "δ=0.75", "圆法兰", "Φ500", "m2", 5, "=1+1", "=G3*H3", None, "二层"])
    sheet.merge_cells("J2:J3")
    sheet.append(["辅材", "角码", "", "", "", "个", 400])
    sheet.append(["", "卡条", "", "", "", "米", 110])
    sheet.append(["不含税合计", "", "", "", "", "", 17])
    buffer = io.BytesIO()
    workbook.save(buffer)
    parsed = read_workbook(buffer.getvalue(), "客户清单.xlsx")[0]
    suggestion = suggest_duct_mapping(parsed)
    assert suggestion["mapping"]["product_state"] == 9
    special = extract_duct_flange_meter_totals(parsed, 2, suggestion["mapping"])
    assert special["totals"] == {"corner_count": "400", "clip_meters": "110"}
    rows = build_quote_rows(parsed, 2, suggestion["mapping"], "duct_v1", special["source_rows"])
    assert len(rows) == 2
    assert [row["attributes"]["product_state"] for row in rows] == ["semi", "semi"]
    assert [row["attributes"]["flange"] for row in rows] == ["common", "round"]
    # Formulas in old quotation-price columns are unrelated to imported inputs.
    assert all(row["source_issue"] is False for row in rows)


def test_round_flange_count_header_is_suggested_without_inference():
    sheet = {"name": "清单", "rows": [
        {"number": 4, "cells": ["构件名称", "规格型号", "法兰方式", "清单总量", "单位", "圆法兰数量"], "issues": []},
        {"number": 5, "cells": ["镀锌风管", "Φ500", "圆法兰", "12", "m2", "8"], "issues": []},
    ]}
    suggestion = suggest_duct_mapping(sheet)
    assert suggestion["mapping"]["round_flange_count"] == 5
    row = build_quote_rows(sheet, 5, suggestion["mapping"], "duct_v1")[0]
    assert row["attributes"]["round_flange_count"] == "8"
    assert row["attributes"]["flange"] == "round"


def test_duct_api_draft_does_not_quote_until_catalog_published(admin):
    quote_id = admin.call("/quotes", "POST", {"pricing_mode": "duct_v1"}).json["id"]
    q = admin.call("/quotes/" + quote_id).json
    assert q["body"]["duct_options"]["tax_included"] is None
    assert q["body"]["duct_options"]["include_common_punch"] is None
    resource_id = admin.call("/resources", "POST", {"kind": "duct_catalog", "body": catalog()["body"]}).json["id"]
    assert q["body"]["duct_options"]["zinc_layer"] == "40"
    options = {"area_basis": "net", "product_state": "semi", "shape": "", "zinc_layer": "40", "flange": "common",
               "corner_count": "", "clip_meters": "",
               "payment_option": "no_install_cash",
               "tax_included": False, "include_common_punch": True, "include_installation": False, "installation_amount": "",
               "include_freight": False}
    row = {"name": "镀锌风管", "spec": "0.5mm", "category": "风管", "engineering_quantity": "100",
           "quantity": "", "unit": "㎡", "remark": "", "attributes": {}, "excluded": False, "exclude_reason": ""}
    edited = admin.call("/quotes/" + quote_id, "PUT", {**q["body"], "rows": [row],
                                                  "duct_options": options, "revision": q["revision"]})
    assert edited.status_code == 200, edited.json
    assert admin.call("/quotes/" + quote_id).json["body"]["duct_options"]["shape"] == "semi_unspecified"
    assert admin.call("/quotes/" + quote_id).json["body"]["duct_options"]["include_common_punch"] is True
    before = admin.call("/quotes/" + quote_id + "/calculate", "POST", {"revision": edited.json["revision"]}).json
    assert before["result"]["rows"][0]["status"] == "pending"
    assert before["result"]["total"] is None
    published = admin.call("/resources/" + resource_id + "/publish", "POST", {})
    assert published.status_code == 200, published.json
    after = admin.call("/quotes/" + quote_id + "/calculate", "POST", {"revision": before["revision"]}).json
    assert after["result"]["rows"][0]["unit_price"] == "25.7"
    assert after["result"]["total"] == "2570"
    assert after["result"]["formal_ready"] is False  # customer template has not been approved


def test_duct_export_separates_flange_and_omits_internal_costs():
    from test_duct_pricing import compute, quote
    q = quote()
    q["duct_options"]["flange"] = "angle"
    q["duct_options"]["include_installation"] = True
    q["duct_options"]["installation_amount"] = "100"
    q["duct_options"]["payment_option"] = "with_install_cash"
    q["rows"][0]["attributes"].update(flange_material="black", flange_size="30",
                                        flange_grade="7.5", flange_meters="10")
    result = compute(q)
    data = write_quote(q, result, "TEST-DUCT", formal=False)
    ws = openpyxl.load_workbook(io.BytesIO(data), data_only=True).active
    assert [ws.cell(4, i).value for i in range(1, 9)] == [
        "序号", "项目名称", "规格型号", "单位", "工程量", "材料单价", "材料合价", "备注"]
    assert [ws.cell(5, i).value for i in (1, 4, 5, 6, 7)] == [1, "㎡", 1000, 19, 19000]
    assert [ws.cell(6, i).value for i in (2, 4, 5, 6, 7)] == ["角钢法兰", "米", 10, 7.2, 72]
    assert "含角钢法兰（按实际米数另列）" in ws["H5"].value
    assert ws["H5"].alignment.wrap_text is True
    assert ws.row_dimensions[5].height >= 28.8
    assert ws.column_dimensions["H"].width == 19.2
    assert ws["A1"].fill.patternType is None
    assert ws.auto_filter.ref is None
    assert "A1:H1" in [str(item) for item in ws.merged_cells.ranges]
    exported = " ".join(str(cell.value or "") for row in ws for cell in row)
    assert "除以0.75" not in exported
    assert "单位成本" not in exported
    assert "不含抗震支架" in exported


def test_duct_export_shows_automatic_freight_amount():
    from test_duct_pricing import compute, quote
    q = quote()
    q["duct_options"]["include_freight"] = True
    result = compute(q)
    ws = openpyxl.load_workbook(io.BytesIO(write_quote(q, result, "TEST-FREIGHT", formal=False)), data_only=True).active
    values = [[cell.value for cell in row] for row in ws]
    freight = next(row for row in values if row[0] == "其中运费（不含税，已计入材料单价）")
    assert freight[6] == 1000
    assert result["total"] == "32400"


def test_duct_export_lists_whole_quote_flange_meters_without_area_conversion():
    from test_duct_pricing import compute, quote
    q = quote()
    q["duct_options"].update(flange_30_meters="10", flange_40_meters="20")
    result = compute(q)
    ws = openpyxl.load_workbook(io.BytesIO(write_quote(q, result, "TEST-AGGREGATE", formal=False)), data_only=True).active
    values = [[cell.value for cell in row] for row in ws]
    assert "整单实际米数另列" in values[4][7]
    flange_30 = next(row for row in values if row[1] == "30#角钢法兰")
    flange_40 = next(row for row in values if row[1] == "40#角钢法兰")
    assert flange_30[3:7] == ["米", 10, 7.2, 72]
    assert flange_40[3:7] == ["米", 20, 9.9, 198]
    assert "替代1.72面积换算" in flange_30[7]


def test_duct_export_lists_round_flange_by_customer_count():
    from test_duct_pricing import compute, quote
    q = quote()
    q["rows"][0]["attributes"].update(flange="round", model="Φ500", round_flange_count="10")
    result = compute(q)
    ws = openpyxl.load_workbook(io.BytesIO(write_quote(q, result, "TEST-ROUND", formal=False)), data_only=True).active
    values = [[cell.value for cell in row] for row in ws]
    round_flange = next(row for row in values if row[1] == "圆法兰")
    assert round_flange[2:7] == ["Φ500mm", "个", 10, 15.7, 157]
    assert "3.14×直径(m)×不含税10元/米" in round_flange[7]
    assert "圆法兰按直径和客户数量另列" in values[4][7]


def test_duct_export_shows_payment_coefficient_and_lists_actual_accessories():
    from test_duct_pricing import compute, quote
    q = quote()
    q["duct_options"].update(
        flange="common", include_common_punch=False, corner_count="100", clip_meters="10",
        payment_option="no_install_days_15",
    )
    result = compute(q)
    ws = openpyxl.load_workbook(
        io.BytesIO(write_quote(q, result, "TEST-PAYMENT", formal=False)), data_only=True).active
    values = [[cell.value for cell in row] for row in ws]
    corner = next(row for row in values if row[1] == "角码")
    clip = next(row for row in values if row[1] == "卡扣")
    assert corner[3:7] == ["个", 100, 0.3, 30]
    assert clip[3:7] == ["米", 10, 3, 30]
    notes = "\n".join(str(row[0] or "") for row in values)
    assert "付款周期7—15天（不含7天，含15天）" in notes
    assert "付款系数1.05" in notes
    assert "付款系数在材料与安装合计后统一应用一次" in notes


def test_duct_export_remarks_follow_current_choices_and_keep_user_note():
    from test_duct_pricing import compute, quote
    q = quote()
    q["duct_options"].update(flange="common", include_common_punch=True, area_basis="net",
                             include_freight=True, tax_included=True)
    q["rows"][0]["remark"] = "客户要求分批交货"
    q["note"] = "有效期另行确认"
    result = compute(q)
    ws = openpyxl.load_workbook(io.BytesIO(write_quote(q, result, "TEST-NOTES", formal=False)), data_only=True).active
    assert "共板法兰" in ws["H5"].value
    assert "含角码卡扣" in ws["H5"].value
    assert "打孔费" in ws["H5"].value
    assert "客户要求分批交货" in ws["H5"].value
    exported = " ".join(str(cell.value or "") for row in ws for cell in row)
    assert "含税" in exported and "含运费" in exported
    assert "有效期另行确认" in exported
    assert "四川中创科庆" not in exported
    assert "7天" not in exported
    assert "2元/平方米" not in exported


def test_formal_duct_export_uses_material_only_template_without_installation():
    from test_duct_pricing import compute, quote
    q = quote()
    q.update(customer="测试客户", project="测试项目", contact="张三 13800000000", delivery_location="贵阳",
             quote_date="2026-09-22", note="客户要求分批交货")
    q["rows"][0]["attributes"]["model"] = "800×320"
    result = compute(q)
    workbook = openpyxl.load_workbook(
        io.BytesIO(write_quote(q, result, "TEST-V1", formal=True)), data_only=True)
    assert workbook.sheetnames == ["单层材料"]
    ws = workbook.active
    assert ws["H1"].value == "KQ-TEST-V1"
    assert ws["B5"].value == "测试客户"
    assert ws["F5"].value == "测试项目"
    assert ws["F6"].value == "贵阳"
    assert ws["B7"].value == "不含税；不含运费；不含安装"
    assert ws["F7"].value == "现金"
    assert ws["G3"].value is None
    assert ws["H3"].value is None
    assert [ws.cell(10, column).value for column in range(1, 11)] == [
        1, "镀锌钢板风管", "0.5", "800×320",
        "角钢法兰", "㎡", 1000, 31.4, 31400, "半成品；不区分直管/异形件；开板面积；含角钢法兰"]
    assert ws["I10"].alignment.horizontal == "center"
    assert len(ws._images) == 1
    assert ws["A11"].value == "价税合计"
    assert ws["I11"].value == 31400
    assert ws["D12"].value == "人民币叁万壹仟肆佰元整"
    assert ws["A14"].value == "报价范围：不含税；不含运费；不含安装。"
    assert "整单付款系数1.0" in ws["A15"].value
    assert "计价口径：开板面积" in ws["A16"].value
    assert "客户要求分批交货" in ws["A18"].value
    exported_text = "\n".join(str(cell.value or "") for row in ws for cell in row)
    assert "本报价有效期为7天" not in exported_text
    assert "标准直管按" not in exported_text
    workbook.close()


def test_formal_duct_export_marks_included_common_accessories_in_remark():
    from test_duct_pricing import compute, quote
    q = quote()
    q["duct_options"].update(flange="common", include_common_punch=False)
    result = compute(q)
    workbook = openpyxl.load_workbook(
        io.BytesIO(write_quote(q, result, "TEST-COMMON", formal=True)), data_only=True)
    remark = workbook.active["J10"].value
    assert "含角码卡扣" in remark
    assert "不含共板法兰打孔费" in remark
    workbook.close()


def test_formal_duct_export_uses_installation_template_and_keeps_installation_separate():
    from test_duct_pricing import compute, quote, set_installation
    q = quote()
    q.update(customer="测试客户", project="测试项目", contact="张三", delivery_location="遵义",
             quote_date="2026-09-22")
    q["duct_options"].update(include_installation=True, payment_option="with_install_cash")
    set_installation(q, '0.1')  # 1000㎡ * .1 = 100; uniform tax produces 101.
    result = compute(q)
    workbook = openpyxl.load_workbook(
        io.BytesIO(write_quote(q, result, "TEST-V2", formal=True)), data_only=True)
    assert workbook.sheetnames == ["带安装"]
    ws = workbook.active
    assert ws["I1"].value == "KQ-TEST-V2"
    assert ws["G6"].value == "遵义"
    assert ws["B7"].value == "不含税；不含运费；含安装（安装统一含税）"
    assert ws["G7"].value == "现金支付（一次性全部付清）"
    assert ws["H3"].value is None and ws["I3"].value is None
    assert ws['F11'].value == '㎡' and ws['G11'].value == 1000
    assert ws['J11'].value == .1 and ws['K11'].value == 100
    assert ws['B12'].value == '安装含税及取整' and ws['K12'].value == 1
    assert ws["I10"].alignment.horizontal == "center"
    assert ws["K11"].alignment.horizontal == "center"
    assert len(ws._images) == 1
    assert ws["A13"].value == "安装小计" and ws['K13'].value == 101
    assert ws["A14"].value == "价税合计" and ws['K14'].value == 31501
    assert ws["D15"].value == "人民币叁万壹仟伍佰零壹元整"
    notes = '\n'.join(str(ws.cell(row, 1).value or '') for row in range(17, 22))
    assert '材料与安装合计后统一乘一次' in notes
    assert '不作用于安装费' not in notes
    for phrase in ('风机控制箱', '抗震支架', '预留洞孔'):
        assert phrase in notes
    workbook.close()


def test_outdated_duct_result_cannot_be_shown_as_current_or_exported(admin):
    quote_id = admin.call("/quotes", "POST", {"pricing_mode": "duct_v1"}).json["id"]
    store = admin.client.application.extensions["store"]
    with store.db() as db:
        db.execute("UPDATE quotes SET result=? WHERE id=?", (dumps({"engine_version": "duct-v2", "total": "9999.0"}), quote_id))
    loaded = admin.call("/quotes/" + quote_id).json
    assert loaded["result"] is None
    assert loaded["result_stale"] is True
    blocked = admin.call("/quotes/" + quote_id + "/export", "POST", {"revision": loaded["revision"], "kind": "draft"})
    assert blocked.status_code == 400
    assert "重新保存并试算" in blocked.json["error"]
    calculated = admin.call("/quotes/" + quote_id + "/calculate", "POST", {"revision": loaded["revision"]})
    assert calculated.json["result"]["engine_version"] == "duct-v16"
    assert admin.call("/quotes/" + quote_id).json["result_stale"] is False
