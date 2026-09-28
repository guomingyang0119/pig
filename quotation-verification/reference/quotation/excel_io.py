import io
import zipfile
import re
from copy import copy
from pathlib import Path
from datetime import datetime, date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from unicodedata import east_asian_width
from math import ceil

import openpyxl
import xlrd
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .engine import RuleError, ATTRIBUTES, text_number

MAX_ROWS, MAX_COLUMNS = 10000, 80
DUCT_ATTRIBUTE_VALUES = {
    "material": {"镀锌板": "galvanized", "40g镀锌板": "galvanized_40", "80g镀锌板": "galvanized_80",
                 "201不锈钢": "stainless_201", "304不锈钢": "stainless_304"},
    "shape": {"直管": "straight", "异形件": "irregular"},
    "flange": {"共板法兰": "common", "共板": "common", "角钢法兰": "angle", "角钢": "angle",
               "角铁法兰": "angle", "角铁": "angle", "圆法兰": "round", "圆形法兰": "round"},
    "flange_material": {"黑角铁": "black", "热镀锌角铁": "hot_galvanized"},
    "flange_size": {"30角钢": "30", "40角钢": "40", "50角钢": "50"},
}


def cell_text(value):
    if value is None: return ""
    if isinstance(value, (date, datetime)): return value.isoformat()
    return str(value)


def read_workbook(content, filename):
    """Read cached values, never execute spreadsheet formulas or macros."""
    extension = Path(filename).suffix.lower()
    sheets = []
    if extension == ".xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(x.file_size for x in archive.infolist()) > 80 * 1024 * 1024:
                    raise RuleError("Excel解压后超过80MB，请拆分清单")
            # Normal mode is intentional: it lets us propagate vertically merged
            # quotation remarks (common in customer workbooks) to every covered row.
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=False)
            formula_wb = openpyxl.load_workbook(io.BytesIO(content), data_only=False, read_only=False)
            try:
                for sheet in wb:
                    rows, issues = [], []
                    if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLUMNS:
                        raise RuleError(f"工作表“{sheet.title}”超过10000行或80列，请拆分")
                    merged_values = {}
                    for merged in sheet.merged_cells.ranges:
                        if merged.min_row >= merged.max_row:
                            continue
                        value = sheet.cell(merged.min_row, merged.min_col).value
                        for merged_row in range(merged.min_row, merged.max_row + 1):
                            for merged_col in range(merged.min_col, merged.max_col + 1):
                                merged_values[(merged_row, merged_col)] = value
                    for index, (cells, formulas) in enumerate(zip(sheet.iter_rows(), formula_wb[sheet.title].iter_rows()), 1):
                        if index > MAX_ROWS or len(cells) > MAX_COLUMNS:
                            raise RuleError("工作表超出导入范围，请拆分；系统未截断导入")
                        values = [cell_text(merged_values.get((index, col + 1), c.value)) for col, c in enumerate(cells)]
                        bad = []
                        for col, (cell, formula) in enumerate(zip(cells, formulas)):
                            if formula.data_type == "f" and cell.value is None:
                                values[col] = "[公式无缓存值]"
                                bad.append(col)
                            elif cell.data_type == "e":
                                bad.append(col)
                        if any(v.strip() for v in values):
                            rows.append({"number": index, "cells": values, "issues": bad})
                    sheets.append({"name": sheet.title, "rows": rows})
            finally:
                wb.close()
                formula_wb.close()
        except RuleError:
            raise
        except Exception:
            raise RuleError("无法读取.xlsx，请确认文件未加密且是有效Excel文件")
    elif extension == ".xls":
        try:
            wb = xlrd.open_workbook(file_contents=content)
            for sheet in wb.sheets():
                if sheet.nrows > MAX_ROWS or sheet.ncols > MAX_COLUMNS:
                    raise RuleError(f"工作表“{sheet.name}”超过导入范围，请拆分")
                rows, merged_values = [], {}
                for row_low, row_high, col_low, col_high in sheet.merged_cells:
                    if row_high - row_low <= 1:
                        continue
                    value = sheet.cell_value(row_low, col_low)
                    for merged_row in range(row_low, row_high):
                        for merged_col in range(col_low, col_high):
                            merged_values[(merged_row, merged_col)] = value
                for index in range(sheet.nrows):
                    values, bad = [], []
                    for col, cell in enumerate(sheet.row(index)):
                        value = merged_values.get((index, col), cell.value)
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate.xldate_as_datetime(value, wb.datemode)
                        if cell.ctype == xlrd.XL_CELL_ERROR: bad.append(col)
                        values.append(cell_text(value))
                    if any(v.strip() for v in values): rows.append({"number": index + 1, "cells": values, "issues": bad})
                sheets.append({"name": sheet.name, "rows": rows})
            wb.release_resources()
        except RuleError:
            raise
        except Exception:
            raise RuleError("无法读取.xls，请确认文件未加密且格式正确")
    else:
        raise RuleError("首期仅支持.xlsx和.xls文件")
    if not any(s["rows"] for s in sheets): raise RuleError("文件没有可读取内容")
    return sheets


def build_quote_rows(sheet, start_row, mapping, pricing_mode="legacy", skip_source_rows=None):
    required = {"name"} if pricing_mode == "duct_v1" else {"name", "quantity", "unit"}
    if any(mapping.get(k) is None for k in required):
        raise RuleError("请映射产品名称列；工程量缺失时可逐行补填" if pricing_mode == "duct_v1" else "请映射产品名称、数量和单位列")
    indexes = [v for v in mapping.values() if v is not None]
    if any(type(v) is not int or not 0 <= v < MAX_COLUMNS for v in indexes): raise RuleError("列映射超出范围")
    if len(set(indexes)) != len(indexes): raise RuleError("不同字段不能映射到同一列")
    if type(start_row) is not int or start_row < 1: raise RuleError("起始行必须是正整数")
    result, skip_source_rows = [], set(skip_source_rows or [])
    for source in sheet["rows"]:
        if source["number"] < start_row or source["number"] in skip_source_rows: continue
        values = source["cells"]
        def get(key):
            col = mapping.get(key)
            return values[col].strip() if col is not None and col < len(values) else ""
        name = get("name")
        category = get("category") or ("风管" if pricing_mode == "duct_v1" and "风管" in name else "")
        def attribute_value(key):
            value = get(key)
            if pricing_mode != "duct_v1":
                return value
            if key == "product_state":
                if "半成品" in value:
                    return "semi"
                if "成品" in value:
                    return "finished"
            return DUCT_ATTRIBUTE_VALUES.get(key, {}).get(value, value)
        mapped_columns = {col for col in mapping.values() if isinstance(col, int)}
        result.append({"source_row": source["number"], "source_sheet": sheet["name"],
                       "original": source, "name": get("name"), "spec": get("spec"),
                       "quantity": get("quantity"), "engineering_quantity": get("engineering_quantity") if pricing_mode == "duct_v1" else "",
                       "unit": get("unit") or ("㎡" if pricing_mode == "duct_v1" else ""), "remark": get("remark"),
                       "category": category,
                       "attributes": {k: attribute_value(k) for k in ATTRIBUTES if get(k)},
                       "excluded": False, "exclude_reason": "",
                       "source_issue": bool(set(source["issues"]) & mapped_columns), "source_issue_confirmed": False})
    if not result: raise RuleError("起始行后没有数据")
    return result


def suggest_duct_mapping(sheet):
    """Preselect only a unique, explicit header; ambiguous columns stay blank."""
    known = {"name": {"产品名称", "材料名称", "项目名称", "构件名称"},
             "engineering_quantity": {"工程量", "总面积", "工程量(㎡)", "工程量（㎡）", "清单总量"},
             "remark": {"备注"}, "unit": {"单位"}, "flange": {"法兰方式", "法兰形式"},
             "product_state": {"报价备注", "成品状态", "风管状态"},
             "round_flange_count": {"圆法兰数量", "圆法兰个数", "法兰数量", "法兰个数"},
             "diameter": {"直径", "直径(mm)", "直径（mm）"}}
    candidates = []
    for row in sheet.get("rows", [])[:20]:
        found = {}
        for key, labels in known.items():
            indexes = [i for i, value in enumerate(row["cells"]) if str(value).strip().replace(" ", "") in labels]
            if len(indexes) == 1:
                found[key] = indexes[0]
            elif len(indexes) > 1:
                found.pop(key, None)
        cells = [str(value).strip().replace(" ", "") for value in row["cells"]]
        descriptive_specs = [i for i, value in enumerate(cells) if value in {"规格范围", "规格描述", "规格"}]
        model_specs = [i for i, value in enumerate(cells) if value == "规格型号"]
        if len(descriptive_specs) == 1:
            found["spec"] = descriptive_specs[0]
            if len(model_specs) == 1:
                found["model"] = model_specs[0]
        elif len(descriptive_specs) == 0 and len(model_specs) == 1:
            found["spec"] = model_specs[0]
        if "name" in found:
            candidates.append((row["number"], found))
    if len(candidates) != 1:
        return {"mapping": {}, "start_row": None}
    header_row, mapping = candidates[0]
    return {"mapping": mapping, "start_row": header_row + 1}


def extract_duct_flange_meter_totals(sheet, start_row, mapping):
    """Extract explicit special rows; never infer meters, accessories or round-flange counts."""
    totals, source_rows, sources = {}, set(), []
    quantity_col, unit_col = mapping.get("engineering_quantity"), mapping.get("unit")
    for source in sheet.get("rows", []):
        if source["number"] < start_row:
            continue
        values = source["cells"]
        text = " ".join(value.strip() for value in values if value.strip())
        normalized = text.replace(" ", "")
        if (normalized.startswith("备注：") or normalized.startswith("备注:") or
                normalized.startswith("报价说明：") or normalized.startswith("报价说明:") or "合计" in normalized):
            source_rows.add(source["number"])
            continue
        unit = values[unit_col].strip().lower() if isinstance(unit_col, int) and unit_col < len(values) else ""
        raw = values[quantity_col].strip() if isinstance(quantity_col, int) and quantity_col < len(values) else ""
        accessory_key = None
        if "角码" in normalized and unit in {"个", "只", "件", "pcs", "pc"}:
            accessory_key = "corner_count"
        elif any(label in normalized for label in ("卡条", "卡码", "卡扣")) and unit in {"米", "m", "meter", "metre"}:
            accessory_key = "clip_meters"
        if accessory_key:
            try:
                quantity = Decimal(raw)
            except (InvalidOperation, ValueError):
                raise RuleError(f"第{source['number']}行辅材数量不是有效数字")
            if not quantity.is_finite() or quantity <= 0:
                raise RuleError(f"第{source['number']}行辅材数量必须大于0")
            if accessory_key in totals:
                raise RuleError("清单存在多行角码或卡扣汇总，请先合并或明确处理")
            totals[accessory_key] = text_number(quantity)
            source_rows.add(source["number"])
            sources.append({"source_row": source["number"], "kind": accessory_key, "quantity": text_number(quantity)})
            continue
        match = re.search(r"(?<!\d)(30|40)\s*#?\s*(?:角钢|角铁)", text)
        if not match or unit not in {"米", "m", "meter", "metre"}:
            continue
        try:
            meters = Decimal(raw)
        except (InvalidOperation, ValueError):
            raise RuleError(f"第{source['number']}行角钢法兰米数不是有效数字")
        if not meters.is_finite() or meters <= 0:
            raise RuleError(f"第{source['number']}行角钢法兰米数必须大于0")
        size = match.group(1)
        key = f"flange_{size}_meters"
        if key in totals:
            raise RuleError(f"清单存在多行{size}#角钢汇总米数，请先合并或明确处理")
        totals[key] = text_number(meters)
        source_rows.add(source["number"])
        sources.append({"source_row": source["number"], "kind": "angle_flange", "size": size, "meters": text_number(meters)})
    return {"totals": totals, "source_rows": source_rows, "sources": sources}


def write_quote(body, result, identifier, formal=False):
    """Application export, not a source workbook mutation. Text stays literal."""
    if body.get("pricing_mode") == "duct_v1":
        if formal:
            return write_duct_formal_quote(body, result, identifier)
        return write_duct_quote(body, result, identifier, formal)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "报价单" if formal else "报价草稿"
    title = "贵州公司 · 报价单" if formal else "贵州公司 · 报价草稿（待核对）"
    ws.merge_cells("A1:I1")
    ws["A1"] = title
    ws["A1"].font = Font(name="微软雅黑", size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="173D39")
    ws.row_dimensions[1].height = 38
    metadata = [f"客户：{body.get('customer', '')}    项目：{body.get('project', '')}",
                f"联系人：{body.get('contact', '')}    报价日期：{body.get('quote_date', '')}",
                f"报价编号：{identifier}    原文件：{body.get('source_filename', '手工录入')}"]
    for i, text in enumerate(metadata, 2):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=9)
        ws.cell(i, 1, text).data_type = "s"
    note = "；".join(result.get("issues", [])) or "请核对产品配置与商务条件"
    ws.merge_cells("A5:I5")
    ws["A5"] = note
    ws["A5"].data_type = "s"
    ws["A5"].alignment = Alignment(wrap_text=True, vertical="center")
    ws["A5"].font = Font(name="微软雅黑", color="A22D2D")
    ws.row_dimensions[5].height = 45
    headers = ["原行", "产品名称", "规格/确认配置", "计价单位", "计费数量", "销售单价", "金额", "状态", "备注"]
    for i, label in enumerate(headers, 1):
        cell = ws.cell(7, i, label)
        cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="235C53")
    ws.row_dimensions[7].height = 26
    for index, (row, value) in enumerate(zip(body["rows"], result["rows"]), 8):
        status = {"priced": "已计算", "pending": "待确认", "excluded": "不计价"}[value["status"]]
        details = row.get("spec", "")
        if value.get("selected_candidate"): details += " / " + value["selected_candidate"]["body"]["name"]
        attrs = row.get("attributes", {})
        public_attrs = [(ATTRIBUTES[k], v) for k, v in attrs.items() if k in ATTRIBUTES and v != ""]
        if public_attrs: details += " / " + "，".join(f"{k}={v}" for k, v in public_attrs)
        values = [str(row.get("source_row", index-7)), row["name"], details,
                  value.get("billing_unit", row.get("unit", "")), value.get("billing_quantity", row.get("quantity", "")),
                  value.get("unit_price") if value["status"] == "priced" else "待确认" if value["status"] == "pending" else "",
                  value.get("amount") if value["status"] == "priced" else "", status,
                  "；".join(x for x in [row.get("remark", ""), value.get("reason", "")] if x)]
        for col, item in enumerate(values, 1):
            cell = ws.cell(index, col, item)
            cell.data_type = "s"
            cell.font = Font(name="微软雅黑", size=10, color="A22D2D" if value["status"] == "pending" else "243B39")
            cell.fill = PatternFill("solid", fgColor="FFF0EF" if value["status"] == "pending" else "F2F7F5" if index % 2 == 0 else "FFFFFF")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = Border(bottom=Side(style="hair", color="D8E1DC"))
        ws.row_dimensions[index].height = 42
    end = 9 + len(body["rows"])
    summary = [("已确定项目小计", result["subtotal"])]
    for key, label in [("installation", "安装费"), ("freight", "运费"), ("tax", "税费")]:
        summary.append((label, result["fees"][key].get("amount") or "待配置/待确认"))
    summary += [("整单取整差额", result.get("rounding_difference") or "待确认"),
                ("最终报价" if formal else "试算总额", result.get("total") or "待确认，当前不是完整报价")]
    for label, value in summary:
        ws.merge_cells(start_row=end, start_column=1, end_row=end, end_column=5)
        ws.cell(end, 1, label).font = Font(name="微软雅黑", bold=True)
        ws.merge_cells(start_row=end, start_column=6, end_row=end, end_column=9)
        ws.cell(end, 6, value).data_type = "s"
        ws.row_dimensions[end].height = 24
        end += 1
    if body.get("note"):
        ws.merge_cells(start_row=end+1, start_column=1, end_row=end+1, end_column=9)
        cell = ws.cell(end+1, 1, "报价备注：" + body["note"])
        cell.data_type = "s"
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[end+1].height = 50
    widths = [9, 26, 40, 12, 16, 19, 21, 14, 48]
    for i, width in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "D8"
    ws.auto_filter.ref = f"A7:I{7+len(body['rows'])}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:7"
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "第 &P 页 / 共 &N 页"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


FORMAL_DUCT_TEMPLATE = (Path(__file__).resolve().parent.parent / "assets" /
                        "正式报价模板_双模板最终精简版_20260922.xlsx")
FORMAL_LOGO = Path(__file__).resolve().parent.parent / "static" / "keqing-logo.png"


def _excel_number(value):
    if value in (None, ""):
        return None
    try:
        parsed = Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError):
        return str(value)
    if parsed == parsed.to_integral_value():
        return int(parsed)
    return float(parsed)


def _set_literal(cell, value):
    cell.value = "" if value is None else str(value)
    cell.data_type = "s"


def _chinese_upper_money(value):
    """Convert the already-confirmed final amount to standard RMB uppercase text."""
    amount = Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    if amount < 0:
        raise RuleError("正式报价总额不能为负数")
    digits = "零壹贰叁肆伍陆柒捌玖"
    group_units = ("", "万", "亿", "兆")

    def four_digits(number):
        parts, pending_zero = [], False
        for divisor, unit in ((1000, "仟"), (100, "佰"), (10, "拾"), (1, "")):
            digit = number // divisor % 10
            if digit:
                if pending_zero and parts:
                    parts.append("零")
                parts.append(digits[digit] + unit)
                pending_zero = False
            elif parts and number % divisor:
                pending_zero = True
        return "".join(parts)

    integer = int(amount)
    groups = []
    remaining = integer
    while remaining:
        groups.append(remaining % 10000)
        remaining //= 10000
    integer_text, pending_zero = [], False
    for index in range(len(groups) - 1, -1, -1):
        group = groups[index]
        if not group:
            if integer_text and any(groups[:index]):
                pending_zero = True
            continue
        if integer_text and (pending_zero or group < 1000):
            integer_text.append("零")
        integer_text.append(four_digits(group) + group_units[index])
        pending_zero = False
    result = "人民币" + ("".join(integer_text) if integer_text else "零") + "元"
    cents = int((amount - integer) * 100)
    jiao, fen = divmod(cents, 10)
    if not jiao and not fen:
        return result + "整"
    if jiao:
        result += digits[jiao] + "角"
    elif fen:
        result += "零"
    if fen:
        result += digits[fen] + "分"
    return result


def _copy_template_row_style(ws, source_row, target_row, last_column):
    for column in range(1, last_column + 1):
        source, target = ws.cell(source_row, column), ws.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        target.protection = copy(source.protection)
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def _resize_formal_detail_area(ws, detail_count, template_capacity=6, summary_row=16):
    """Resize rows 10:15 while preserving all template merges below the detail area."""
    if detail_count < 1:
        raise RuleError("正式报价没有可输出的明细")
    delta = detail_count - template_capacity
    if not delta:
        return
    shifted = []
    for merged in list(ws.merged_cells.ranges):
        if merged.min_row >= summary_row:
            shifted.append((str(merged), merged.min_row, merged.max_row, merged.min_col, merged.max_col))
            ws.unmerge_cells(str(merged))
    if delta > 0:
        ws.insert_rows(summary_row, delta)
    else:
        ws.delete_rows(summary_row + delta, -delta)
    for _, min_row, max_row, min_col, max_col in shifted:
        ws.merge_cells(start_row=min_row + delta, start_column=min_col,
                       end_row=max_row + delta, end_column=max_col)


def _formal_specification(source):
    """Return the product size/model only; thickness has its own output column."""
    attributes = source.get("attributes") or {}
    model = str(attributes.get("model") or "").strip()
    if model:
        return model
    specification = str(source.get("spec") or "").strip()
    if not specification:
        return ""
    rectangle = re.search(
        r"(?<!\d)(\d+(?:\.\d+)?)\s*[×xX*＊]\s*(\d+(?:\.\d+)?)(?:\s*(?:mm|毫米))?",
        specification)
    if rectangle:
        return f"{rectangle.group(1)}×{rectangle.group(2)}"
    round_size = re.search(r"(?:Φ|φ|Ø|ø)\s*\d+(?:\.\d+)?(?:\s*(?:mm|毫米))?", specification)
    if round_size:
        return re.sub(r"\s+", "", round_size.group(0))
    dn_size = re.search(r"\bDN\s*\d+(?:\.\d+)?\b", specification, flags=re.IGNORECASE)
    if dn_size:
        return re.sub(r"\s+", "", dn_size.group(0)).upper()
    without_thickness = re.sub(
        r"(?:δ\s*=?\s*|厚度\s*[:：]?\s*)\d+(?:\.\d+)?\s*(?:mm|毫米)?",
        "", specification, flags=re.IGNORECASE)
    without_thickness = re.sub(
        r"^\s*\d+(?:\.\d+)?\s*(?:mm|毫米)?\s*(?:[/／,，;；|｜-]+\s*)?",
        "", without_thickness, flags=re.IGNORECASE)
    return without_thickness.strip(" /／,，;；|｜-")


def _formal_duct_lines(body, result, include_installation):
    options = body.get("duct_options", {})
    area_label = "净面积" if options.get("area_basis") == "net" else "开板面积"
    flange_labels = {"common": "共板法兰", "angle": "角钢法兰", "round": "圆法兰"}
    state_labels = {"semi": "半成品", "finished": "成品"}
    shape_labels = {"semi_unspecified": "不区分直管/异形件", "straight": "直管", "irregular": "异形件"}
    lines = []

    def line(**values):
        base = {"name": "", "thickness": "", "spec": "", "process": "", "unit": "",
                "quantity": None, "material_unit": None, "material_amount": None,
                "installation_unit": None, "installation_amount": None, "remark": ""}
        base.update(values)
        lines.append(base)

    for source, calculated in zip(body.get("rows", []), result.get("rows", [])):
        if calculated.get("status") != "priced":
            continue
        selected_material = calculated.get("selected_material") or {}
        flange = flange_labels.get(calculated.get("effective_flange"), "法兰形式已确认")
        state = state_labels.get(calculated.get("effective_product_state"), "")
        shape = shape_labels.get(calculated.get("selected_shape"), "")
        process = flange
        specification = _formal_specification(source)
        remarks = [item for item in (state, shape, area_label) if item]
        basis = calculated.get("flange_basis")
        if basis == "common_punch":
            remarks.append("含共板法兰打孔费")
        elif basis == "common_no_punch":
            remarks.append("不含共板法兰打孔费")
        elif basis == "actual_meters":
            remarks.append("角钢法兰按实际米数另列")
        elif basis == "aggregate_meters":
            remarks.append("角钢法兰按整单实际米数另列")
        elif basis == "round_count":
            remarks.append("圆法兰按客户数量另列")
        elif basis == "round_area":
            remarks.append("圆法兰按面积折算")
        if source.get("remark"):
            remarks.append(str(source["remark"]))
        line(name=source.get("name", ""), thickness=selected_material.get("thickness", ""),
             spec=specification, process=process,
             unit=calculated.get("billing_unit") or source.get("unit") or "㎡",
             quantity=_excel_number(calculated.get("billing_quantity")),
             material_unit=_excel_number(calculated.get("unit_price")),
             material_amount=_excel_number(calculated.get("body_amount")),
             remark="；".join(remarks))
        if calculated.get("flange_amount") is not None:
            selected_flange = calculated.get("selected_flange") or {}
            line(name="角钢法兰", thickness="-", spec=selected_flange.get("name", ""),
                 process="角钢法兰", unit="米",
                 quantity=_excel_number(calculated.get("flange_meters")),
                 material_unit=_excel_number(calculated.get("flange_unit_price")),
                 material_amount=_excel_number(calculated.get("flange_amount")),
                 remark="按客户清单实际米数计价；不再乘1.72")
        if calculated.get("round_flange_amount") is not None:
            line(name="圆法兰", thickness="-",
                 spec=f"Φ{calculated.get('round_flange_diameter')}mm", process="圆法兰", unit="个",
                 quantity=_excel_number(calculated.get("round_flange_count")),
                 material_unit=_excel_number(calculated.get("round_flange_unit_price")),
                 material_amount=_excel_number(calculated.get("round_flange_amount")),
                 remark="按客户清单数量计价；圆法兰数量由客户清单提供")

    for accessory in result.get("accessory_items", []):
        line(name=accessory.get("name", "整单配件"), thickness="-", spec="按客户清单实际数量",
             process="共板法兰配件", unit=accessory.get("unit", ""),
             quantity=_excel_number(accessory.get("quantity")),
             material_unit=_excel_number(accessory.get("unit_price")),
             material_amount=_excel_number(accessory.get("amount")), remark="整单实际数量")
    for flange in result.get("aggregate_flanges", []):
        line(name=f"{flange['size']}#角钢法兰", thickness="-", spec=flange.get("name", ""),
             process="角钢法兰", unit="米", quantity=_excel_number(flange.get("meters")),
             material_unit=_excel_number(flange.get("unit_price")),
             material_amount=_excel_number(flange.get("amount")),
             remark="按整单实际米数计价；不再乘1.72面积换算")
    if include_installation:
        installation = result.get("fees", {}).get("installation", {})
        amount = _excel_number(installation.get("amount"))
        line(name="安装费", thickness="-", spec="整单安装费", process="安装", unit="项",
             quantity=1, installation_unit=amount, installation_amount=amount,
             remark="报价人员确认金额；安装费不乘付款系数；" +
                    ("含税" if options.get("tax_included") is True else "不含税"))
    return lines


def _formal_quote_notes(body, result, include_installation):
    """Build only remarks confirmed by the current quotation; template prose is not reused."""
    options = body.get("duct_options", {})
    scope = ["含税" if options.get("tax_included") is True else "不含税"]
    scope.append("含运费" if options.get("include_freight") is True else "不含运费")
    scope.append("含安装" if include_installation else "不含安装")
    notes = ["报价范围：" + "；".join(scope) + "。"]

    payment = result.get("payment") or {}
    if payment.get("label"):
        payment_note = "付款条件：" + str(payment["label"])
        if payment.get("coefficient") not in (None, ""):
            payment_note += "；材料付款系数" + str(payment["coefficient"])
        payment_note += "，作用于供货材料单价"
        if options.get("include_freight") is True:
            payment_note += "（含运费）"
        if include_installation:
            payment_note += "，不作用于安装费"
        notes.append(payment_note + "。")

    area_labels = {"open": "开板面积", "net": "净面积"}
    area_label = area_labels.get(options.get("area_basis"))
    if area_label:
        notes.append(
            f"计价口径：{area_label}；材料规格、成品状态、管件类型、法兰形式及特殊费用以报价明细为准。")
    notes.append("实际供货数量、厚度、规格型号及工艺，以双方确认的图纸、清单或订单为准。")

    if include_installation:
        notes.append(
            "安装报价排除范围：①不含风机控制箱及控制箱到风机的电线、电器；"
            "②不含风机基础预制的制作及抗震支架；③不含设备及管道预留洞孔的开启与修复。")
    if body.get("note"):
        quote_note = "报价备注：" + str(body["note"])
        if len(notes) < 5:
            notes.append(quote_note)
        else:
            notes[-1] += "\n" + quote_note
    return notes[:5]


def _write_formal_quote_notes(ws, first_row, notes):
    """Replace the five template note rows with current-quotation remarks."""
    for offset in range(5):
        row_number = first_row + offset
        value = notes[offset] if offset < len(notes) else ""
        cell = ws.cell(row_number, 1)
        _set_literal(cell, value)
        alignment = copy(cell.alignment)
        alignment.horizontal = "left"
        alignment.vertical = "center"
        alignment.wrap_text = True
        cell.alignment = alignment
        if value:
            estimated_lines = max(1, ceil(sum(2 if east_asian_width(char) in "FW" else 1
                                              for char in value) / 105))
            ws.row_dimensions[row_number].height = max(31, estimated_lines * 18)
        else:
            ws.row_dimensions[row_number].height = 18


def write_duct_formal_quote(body, result, identifier):
    """Fill the approved material-only or material-plus-installation template."""
    if not FORMAL_DUCT_TEMPLATE.is_file():
        raise RuleError("正式报价双模板文件缺失，不能生成正式报价")
    if not FORMAL_LOGO.is_file():
        raise RuleError("公司Logo文件缺失，不能生成正式报价")
    include_installation = body.get("duct_options", {}).get("include_installation") is True
    sheet_name = "带安装" if include_installation else "单层材料"
    try:
        wb = openpyxl.load_workbook(FORMAL_DUCT_TEMPLATE)
        if sheet_name not in wb.sheetnames:
            raise RuleError(f"正式报价模板缺少“{sheet_name}”工作表")
        ws = wb[sheet_name]
        for other in list(wb.worksheets):
            if other is not ws:
                wb.remove(other)
        ws._images.clear()
        logo = XLImage(FORMAL_LOGO)
        logo.width, logo.height = 118, 104
        ws.add_image(logo, "A1")
        lines = _formal_duct_lines(body, result, include_installation)
        _resize_formal_detail_area(ws, len(lines))
        last_column = 12 if include_installation else 10
        for row_number, item in enumerate(lines, 10):
            _copy_template_row_style(ws, 10, row_number, last_column)
            values = [row_number - 9, item["name"], item["thickness"], item["spec"], item["process"],
                      item["unit"], item["quantity"], item["material_unit"], item["material_amount"]]
            if include_installation:
                values.extend([item["installation_unit"], item["installation_amount"], item["remark"]])
            else:
                values.append(item["remark"])
            for column, value in enumerate(values, 1):
                cell = ws.cell(row_number, column)
                if isinstance(value, str):
                    _set_literal(cell, value)
                else:
                    cell.value = value
                if column in (8, 9, 10, 11) and isinstance(value, (int, float)):
                    cell.number_format = "0.0"
                if column == 9 or (include_installation and column == 11):
                    alignment = copy(cell.alignment)
                    alignment.horizontal = "center"
                    cell.alignment = alignment
            ws.row_dimensions[row_number].height = 34

        options = body.get("duct_options", {})
        tax = "含税" if options.get("tax_included") is True else "不含税"
        freight = "含运费" if options.get("include_freight") is True else "不含运费"
        installation = "含安装" if include_installation else "不含安装"
        scope = f"{tax}；{freight}；{installation}"
        payment = (result.get("payment") or {}).get("label") or ""
        quote_number = "KQ-" + str(identifier)
        if include_installation:
            metadata = {"I1": quote_number, "I2": body.get("quote_date", ""),
                        "B5": body.get("customer", ""), "G5": body.get("project", ""),
                        "B6": body.get("contact", ""), "G6": body.get("delivery_location", ""),
                        "B7": scope, "G7": payment}
        else:
            metadata = {"H1": quote_number, "H2": body.get("quote_date", ""),
                        "B5": body.get("customer", ""), "F5": body.get("project", ""),
                        "B6": body.get("contact", ""), "F6": body.get("delivery_location", ""),
                        "B7": scope, "F7": payment}
        for coordinate, value in metadata.items():
            _set_literal(ws[coordinate], value)
        for coordinate in (("H3", "I3") if include_installation else ("G3", "H3")):
            _set_literal(ws[coordinate], "")

        detail_end = 9 + len(lines)
        total = _excel_number(result.get("total"))
        if total is None:
            raise RuleError("正式报价总额尚未确认")
        if include_installation:
            installation_row, total_row, uppercase_row = detail_end + 1, detail_end + 2, detail_end + 3
            ws[f"K{installation_row}"] = _excel_number(result["fees"]["installation"]["amount"])
            ws[f"K{total_row}"] = total
            _set_literal(ws[f"D{uppercase_row}"], _chinese_upper_money(result["total"]))
            for row, height in ((installation_row, 27), (total_row, 31), (uppercase_row, 28),
                                (uppercase_row + 1, 25), (uppercase_row + 2, 31),
                                (uppercase_row + 3, 31), (uppercase_row + 4, 31),
                                (uppercase_row + 5, 45), (uppercase_row + 6, 36)):
                ws.row_dimensions[row].height = height
            first_note_row, last_note_row = uppercase_row + 2, uppercase_row + 6
            ws[f"K{installation_row}"].number_format = "0.0"
            ws[f"K{total_row}"].number_format = "0.0"
        else:
            total_row, uppercase_row = detail_end + 1, detail_end + 2
            ws[f"I{total_row}"] = total
            _set_literal(ws[f"D{uppercase_row}"], _chinese_upper_money(result["total"]))
            for row, height in ((total_row, 31), (uppercase_row, 28), (uppercase_row + 1, 25),
                                (uppercase_row + 2, 31), (uppercase_row + 3, 31),
                                (uppercase_row + 4, 31), (uppercase_row + 5, 45),
                                (uppercase_row + 6, 31)):
                ws.row_dimensions[row].height = height
            first_note_row, last_note_row = uppercase_row + 2, uppercase_row + 6
            ws[f"I{total_row}"].number_format = "0.0"
        _write_formal_quote_notes(
            ws, first_note_row, _formal_quote_notes(body, result, include_installation))
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1 if len(lines) <= 12 else 0
        ws.print_title_rows = "1:9"
        ws.print_area = f"A1:{get_column_letter(last_column)}{last_note_row}"
        ws.print_options.horizontalCentered = True
        out = io.BytesIO()
        wb.save(out)
        wb.close()
        return out.getvalue()
    except RuleError:
        raise
    except Exception as error:
        raise RuleError("正式报价模板读取或填充失败，请检查模板文件") from error


def write_duct_quote(body, result, identifier, formal=False):
    """Reproduce the supplied eight-column quotation layout with current duct data."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "风管、风阀、风口"
    grid = Side(style="thin", color="777777")
    box = Border(left=grid, right=grid, top=grid, bottom=grid)
    for index, width in enumerate([6.4, 28.6, 43.9, 8.8, 9.2, 10.1, 12.9, 19.2], 1):
        ws.column_dimensions[get_column_letter(index)].width = width

    def put(row, col, value, *, size=9, bold=False, horizontal="center", wrap=True, border=True):
        cell = ws.cell(row, col, value if value is not None else "")
        if isinstance(value, str):
            cell.data_type = "s"  # Source text must not become an Excel formula.
        cell.font = Font(name="宋体", size=size, bold=bold, color="A22D2D" if value == "待确认" else "202020")
        cell.alignment = Alignment(horizontal=horizontal, vertical="center", wrap_text=wrap)
        if border:
            cell.border = box
        return cell

    def display_width(value):
        return sum(2 if east_asian_width(char) in "FW" else 1 for char in str(value))

    def merged_box(row, first, last):
        for col in range(first, last + 1):
            ws.cell(row, col).border = box

    ws.merge_cells("A1:H1")
    heading = put(1, 1, "贵州公司" if formal else "贵州公司（报价草稿·待核对）", size=24, border=False)
    heading.font = Font(name="微软雅黑", size=24, color="202020")
    ws.row_dimensions[1].height = 102
    ws.merge_cells("A2:C2")
    ws.merge_cells("D2:E2")
    ws.merge_cells("F2:G2")
    put(2, 1, f"项目：{body.get('project') or '—'}    客户：{body.get('customer') or '—'}", size=11, horizontal="left")
    put(2, 4, f"联系人：{body.get('contact') or '—'}", size=11, horizontal="left")
    put(2, 6, f"编号：{identifier}", size=11, horizontal="left")
    put(2, 8, body.get("quote_date") or "", size=11)
    merged_box(2, 1, 8)
    ws.row_dimensions[2].height = 30
    ws.merge_cells("A3:H3")
    put(3, 1, "风管报价单", size=12)
    merged_box(3, 1, 8)
    ws.row_dimensions[3].height = 30
    options = body.get("duct_options", {})
    flags = ["含税" if options.get("tax_included") is True else "不含税" if options.get("tax_included") is False else "税费待确认",
             "净面积" if options.get("area_basis") == "net" else "开板面积" if options.get("area_basis") == "open" else "面积口径待确认",
             "半成品" if options.get("product_state") == "semi" else "成品" if options.get("product_state") == "finished" else "成品状态待确认",
             "角钢法兰" if options.get("flange") == "angle" else "共板法兰" if options.get("flange") == "common" else "圆法兰" if options.get("flange") == "round" else "混合法兰" if options.get("flange") == "mixed" else "法兰待确认"]
    payment = result.get("payment") or {}
    payment_summary = payment.get("label") or "付款方式待确认"
    if payment.get("coefficient"):
        payment_summary += f"；付款系数{payment['coefficient']}"
    headers = ["序号", "项目名称", "规格型号", "单位", "工程量", "材料单价", "材料合价", "备注"]
    for index, label in enumerate(headers, 1):
        put(4, index, label, size=12)
    ws.row_dimensions[4].height = 28.8

    def number(value):
        if value is None or value == "":
            return None
        try:
            return float(Decimal(str(value)))
        except (InvalidOperation, ValueError, TypeError):
            return str(value)

    def row_remark(row, calculated):
        if calculated["status"] == "pending":
            parts = ["待确认：" + (calculated.get("reason") or "报价条件未齐")]
        elif calculated["status"] == "excluded":
            parts = ["不计价：" + (row.get("exclude_reason") or "已排除")]
        else:
            effective_flange = {"common": "共板法兰", "angle": "角钢法兰", "round": "圆法兰"}.get(calculated.get("effective_flange"), flags[3])
            effective_state = {"semi": "半成品", "finished": "成品"}.get(calculated.get("effective_product_state"), flags[2])
            parts = [f"{effective_flange}{effective_state}（{flags[1]}）"]
            basis = calculated.get("flange_basis")
            if basis == "actual_meters":
                parts.append("法兰按米另列")
            elif basis == "aggregate_meters":
                parts.append("法兰按整单实际米数另列")
            elif basis == "round_count":
                parts.append("圆法兰按直径和客户数量另列")
            elif basis == "round_area":
                parts.append("圆法兰无数量，按直径×材料价×1.72折算到㎡")
            elif basis == "per_area":
                parts.append("含法兰")
            elif basis == "common_punch":
                parts.append("含打孔费")
            elif basis == "common_no_punch":
                parts.append("不含打孔费")
        if row.get("remark"):
            parts.append(row["remark"])
        return "；".join(parts)

    def append_line(items, pending=False):
        index = ws.max_row + 1
        for col, item in enumerate(items, 1):
            cell = put(index, col, item, horizontal="left" if col == 2 else "center")
            if col in (6, 7) and isinstance(item, (int, float)):
                cell.number_format = "0.0"
            if pending:
                cell.font = Font(name="宋体", size=9, color="A22D2D")
        remark_lines = ceil(display_width(items[7]) / 24) + (1 if "；" in str(items[7]) and display_width(items[7]) > 55 else 0)
        lines = max(ceil(display_width(items[1]) / 38), ceil(display_width(items[2]) / 60), remark_lines)
        ws.row_dimensions[index].height = max(28.8, 15 * max(1, lines) + 2)

    item_count = 0
    for row, calculated in zip(body["rows"], result["rows"]):
        item_count += 1
        status = calculated["status"]
        model = str(row.get("attributes", {}).get("model") or "").strip()
        specification = row.get("spec", "")
        if model and model not in specification:
            specification = (specification + " / " + model).strip(" /")
        line = [item_count, row.get("name", ""), specification, row.get("unit") or "㎡",
                number(calculated.get("billing_quantity", row.get("engineering_quantity", "")) if status == "priced" else row.get("engineering_quantity", "")),
                number(calculated.get("unit_price")) if status == "priced" else "待确认" if status == "pending" else "",
                number(calculated.get("body_amount")) if status == "priced" else "",
                row_remark(row, calculated)]
        append_line(line, status == "pending")
        if status == "priced" and calculated.get("flange_amount") is not None:
            flange = calculated.get("selected_flange") or {}
            append_line(["", "角钢法兰", flange.get("name", ""), "米", number(calculated.get("flange_meters")),
                         number(calculated.get("flange_unit_price")), number(calculated.get("flange_amount")),
                         f"按实际米数另计；已乘付款系数{payment.get('coefficient') or '待确认'}"])
        if status == "priced" and calculated.get("round_flange_amount") is not None:
            tax_label = "含税12元/米" if options.get("tax_included") is True else "不含税10元/米"
            append_line(["", "圆法兰", f"Φ{calculated.get('round_flange_diameter')}mm", "个",
                         number(calculated.get("round_flange_count")), number(calculated.get("round_flange_unit_price")),
                         number(calculated.get("round_flange_amount")), f"3.14×直径(m)×{tax_label}×付款系数{payment.get('coefficient') or '待确认'}；按客户清单数量"])

    for accessory in result.get("accessory_items", []):
        item_count += 1
        append_line([item_count, accessory.get("name", "整单配件"), "按客户清单实际数量", accessory.get("unit", ""),
                     number(accessory.get("quantity")), number(accessory.get("unit_price")), number(accessory.get("amount")),
                     f"基础单价×付款系数{payment.get('coefficient') or '待确认'}" + ("×1.1含税系数" if options.get("tax_included") is True else "（不含税）")])

    for flange in result.get("aggregate_flanges", []):
        item_count += 1
        append_line([item_count, f"{flange['size']}#角钢法兰", flange.get("name", ""), "米",
                     number(flange.get("meters")), number(flange.get("unit_price")), number(flange.get("amount")),
                     f"客户清单整单实际米数；替代1.72面积换算；已乘付款系数{payment.get('coefficient') or '待确认'}"])

    def summary(label, value):
        index = ws.max_row + 1
        for col in range(1, 9):
            put(index, col, "")
        ws.merge_cells(start_row=index, start_column=1, end_row=index, end_column=3)
        put(index, 1, label, size=12 if label in ("合计", "试算合计") else 9)
        merged_box(index, 1, 3)
        display = number(value) if value is not None else "待确认"
        cell = put(index, 7, display, size=12 if label in ("合计", "试算合计") else 9)
        if isinstance(display, (int, float)):
            cell.number_format = "0.0"
        ws.row_dimensions[index].height = 30

    has_fees = (result.get("accessory_mode") == "actual" or options.get("include_freight") or
                options.get("include_installation") or options.get("tax_included") is True)
    if has_fees:
        summary("材料小计（已确定项目）", result["subtotal"])
    if options.get("include_freight"):
        summary("其中运费（含税，已计入材料单价）" if options.get("tax_included") is True else "其中运费（不含税，已计入材料单价）", result["fees"]["freight"]["amount"])
    if options.get("include_installation"):
        summary("安装费（含税，已计入）" if options.get("tax_included") is True else "安装费（不含税）", result["fees"]["installation"]["amount"])
    if options.get("tax_included") is True:
        summary("其中税费/含税差额（已计入单价及费用，不重复加总）", result["fees"]["tax"]["amount"])
    summary("合计" if formal else "试算合计", result["total"] if result["total"] is not None else "待确认，非完整报价")
    conditions = "、".join(flags)
    notes = ["报价说明：", "① 本单按" + conditions + "报价；" +
              ("含运费" if options.get("include_freight") else "不含运费") + "；" +
              ("含安装费" if options.get("include_installation") else "不含安装费") + "。",
             "② 付款条件：" + payment_summary + "。付款系数应用于全部供货材料（包括运费），不应用于安装费。"]
    if result.get("issues"):
        notes.append("待确认事项：" + "；".join(result["issues"]))
    for note in result.get("installation_notes", []):
        notes.append("安装范围说明：" + note)
    if body.get("note"):
        notes.append("报价备注：" + body["note"])
    note_row = ws.max_row + 1
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=8)
    put(note_row, 1, "\n".join(notes), size=11, horizontal="left")
    merged_box(note_row, 1, 8)
    note_lines = sum(max(1, ceil(display_width(line) / 115)) for line in notes)
    ws.row_dimensions[note_row].height = min(409, max(50, 17 * note_lines + 8))
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1 if len(body["rows"]) <= 20 else 0
    ws.print_title_rows = "1:4"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
