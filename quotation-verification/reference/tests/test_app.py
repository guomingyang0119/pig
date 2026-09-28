import copy
import io
import json
from pathlib import Path

import openpyxl
import pytest

from conftest import ApiClient
from quotation.excel_io import read_workbook, build_quote_rows


def workbook_bytes():
    wb=openpyxl.Workbook()
    ws=wb.active
    ws.title="客户清单"
    ws.append(["材料清单"])
    ws.append(["产品名称","规格","数量","单位"])
    ws.append(["测试风口","400x400",2,"个"])
    ws.append(["合计",None,None,None])
    stream=io.BytesIO();wb.save(stream)
    return stream.getvalue()


def new_quote(admin, product):
    created=admin.call('/quotes','POST',{'project':'测试项目','customer':'测试客户'}).json
    identifier=created['id']
    current=admin.call('/quotes/'+identifier).json
    response=admin.call('/quotes/'+identifier,'PUT',{**current['body'],'rows':[product],'revision':current['revision']})
    assert response.status_code==200,response.json
    return admin.call('/quotes/'+identifier).json


def publish_rule(admin, rule, product):
    response=admin.call('/resources','POST',{'kind':'rule','body':rule['body']})
    identifier=response.json['id']
    response=admin.call('/resources/'+identifier+'/publish','POST',{'sample':product,'expected_amount':'200'})
    assert response.status_code==200,response.json
    return identifier


def set_policy(admin, policy):
    p=copy.deepcopy(policy)
    p['version']=admin.call('/settings').json['policy']['version']
    response=admin.call('/settings','PUT',{'policy':p})
    assert response.status_code==200,response.json


def test_setup_and_cross_origin_protection(app, admin):
    assert admin.call('/setup','POST',{'username':'x','password':'12345678'}).status_code==409
    assert admin.call('/quotes','POST',{},headers={'Origin':'https://untrusted.example'}).status_code==403
    assert admin.call('/quotes','POST',{},headers={'X-CSRF-Token':'wrong'}).status_code==403
    assert admin.call('/health',headers={'Host':'untrusted.example'}).status_code==403
    response=admin.call('/bootstrap')
    assert 'password_hash' not in json.dumps(response.json)
    assert response.json['user']['role']=='leader'


def test_quoter_permissions_and_ownership(app, admin, product):
    own=new_quote(admin,product)
    response=admin.call('/users','POST',{'username':'staff','display_name':'测试报价员','password':'Test-staff-password','role':'quoter'})
    assert response.status_code==200
    user=ApiClient(app);user.login('staff','Test-staff-password')
    assert user.call('/resources','POST',{'kind':'rule','body':{}}).status_code==403
    assert user.call('/settings','PUT',{}).status_code==403
    assert user.call('/backup','POST',{}).status_code==403
    assert user.call('/quotes/'+own['id']).status_code==404
    assert user.call('/quotes').json==[]
    q=user.call('/quotes','POST',{'project':'员工项目'}).json
    assert admin.call('/quotes/'+q['id']).status_code==200


def test_upload_mapping_preserves_every_nonempty_data_row(admin):
    data={'file':(io.BytesIO(workbook_bytes()),'测试.xlsx')}
    response=admin.client.post('/api/imports',data=data,headers={'X-CSRF-Token':admin.csrf},content_type='multipart/form-data')
    assert response.status_code==200,response.json
    imp=response.json
    response=admin.call('/quotes','POST',{'import_id':imp['id'],'sheet':'客户清单','start_row':3,'mapping':{'name':0,'spec':1,'quantity':2,'unit':3}})
    q=admin.call('/quotes/'+response.json['id']).json
    assert len(q['body']['rows'])==2
    assert q['body']['rows'][1]['name']=='合计'
    assert q['body']['rows'][0]['source_row']==3
    response=admin.call('/quotes/'+q['id'],'PUT',{**q['body'],'revision':q['revision'],'rows':q['body']['rows'][:1]})
    assert response.status_code==400


def test_uncached_formula_is_not_treated_as_blank():
    wb=openpyxl.Workbook();wb.active.append(['产品','=2+3','台'])
    data=io.BytesIO();wb.save(data)
    parsed=read_workbook(data.getvalue(),'formula.xlsx')
    rows=build_quote_rows(parsed[0],1,{'name':0,'quantity':1,'unit':2})
    assert rows[0]['source_issue']
    assert rows[0]['quantity']=='[公式无缓存值]'


def test_pending_export_red_no_internal_cost_and_no_formula_injection(admin, product):
    product['name']='=HYPERLINK("https://example.com")'
    product['attributes']={'unit_cost':'987.654321','margin':'0.25','material':'测试材质'}
    q=new_quote(admin,product)
    response=admin.call('/quotes/'+q['id']+'/calculate','POST',{'revision':q['revision']})
    assert response.json['result']['rows'][0]['amount'] is None
    response=admin.call('/quotes/'+q['id']+'/export','POST',{'revision':response.json['revision'],'kind':'draft'})
    assert response.status_code==200,response.json
    content=admin.client.get(response.json['download']).data
    wb=openpyxl.load_workbook(io.BytesIO(content),data_only=False)
    ws=wb.active
    assert ws['B8'].data_type=='s' and ws['B8'].value.startswith('=HYPERLINK')
    assert ws['F8'].value=='待确认'
    assert ws['B8'].fill.fgColor.rgb=='00FFF0EF'
    texts=' '.join(str(c.value) for row in ws for c in row if c.value is not None)
    assert '单位成本' not in texts and '利润' not in texts
    assert '987.654321' not in texts and 'unit_cost' not in texts and 'margin' not in texts
    assert '测试材质' in texts
    q=admin.call('/quotes/'+q['id']).json
    assert admin.call('/quotes/'+q['id']+'/export','POST',{'revision':q['revision'],'kind':'formal'}).status_code==400


def test_complete_quote_immutable_history_and_new_rule_version(admin, product, rule, policy):
    set_policy(admin,policy)
    rid=publish_rule(admin,rule,product)
    q=new_quote(admin,product)
    response=admin.call('/quotes/'+q['id']+'/calculate','POST',{'revision':q['revision']})
    assert response.json['result']['total']=='200'
    assert response.json['result']['formal_ready']
    response=admin.call('/quotes/'+q['id']+'/export','POST',{'revision':response.json['revision'],'kind':'formal'})
    assert response.status_code==200,response.json
    q=admin.call('/quotes/'+q['id']).json
    assert q['state']=='final'
    assert admin.call('/quotes/'+q['id'],'PUT',{**q['body'],'revision':q['revision']}).status_code==409
    assert admin.call('/resources/'+rid,'PUT',{'body':rule['body']}).status_code==409
    rev=admin.call('/resources/'+rid+'/revise','POST',{}).json['id']
    modified=copy.deepcopy(rule['body']);modified['parameters']['unit_cost']='90'
    assert admin.call('/resources/'+rev,'PUT',{'body':modified}).status_code==200
    assert admin.call('/resources/'+rev+'/publish','POST',{'sample':product,'expected_amount':'240'}).status_code==200
    assert admin.call('/quotes/'+q['id']).json['result']['total']=='200'
    revised=admin.call('/quotes/'+q['id']+'/revise','POST',{}).json['id']
    q2=admin.call('/quotes/'+revised).json
    assert q2['version']==2
    computed=admin.call('/quotes/'+revised+'/calculate','POST',{'revision':q2['revision']}).json
    assert computed['result']['total']=='240'


def test_rule_publish_requires_independent_expected_amount(admin,rule,product):
    rid=admin.call('/resources','POST',{'kind':'rule','body':rule['body']}).json['id']
    assert admin.call('/resources/'+rid+'/publish','POST',{'sample':product,'expected_amount':'201'}).status_code==400
    resources=admin.call('/resources').json
    assert resources[0]['status']=='draft'


def test_quote_optimistic_concurrency_and_original_values(admin,product):
    q=new_quote(admin,product)
    response=admin.call('/quotes/'+q['id'],'PUT',{**q['body'],'revision':q['revision'],'result':{'total':'999'}})
    assert response.status_code==200
    assert admin.call('/quotes/'+q['id']).json['result'] is None
    assert admin.call('/quotes/'+q['id'],'PUT',{**q['body'],'revision':q['revision']}).status_code==409


def test_backup_restore_round_trip_and_sessions_invalidated(admin,app,product):
    original=new_quote(admin,product)
    response=admin.call('/backup','POST',{})
    backup=admin.client.get(response.json['download']).data
    new_quote(admin,product)
    assert len(admin.call('/quotes').json)==2
    response=admin.client.post('/api/restore',data={'file':(io.BytesIO(backup),'backup.sqlite3'),'confirmation':'恢复备份'},headers={'X-CSRF-Token':admin.csrf},content_type='multipart/form-data')
    assert response.status_code==200,response.json
    assert admin.call('/quotes').status_code==401
    admin.login('leader','Test-password-291!')
    quotes=admin.call('/quotes').json
    assert len(quotes)==1 and quotes[0]['id']==original['id']
    assert list((app.extensions['store'].directory/'backups').glob('before_restore_*.sqlite3'))


def test_reject_foreign_backup_without_replacing_database(admin):
    response=admin.client.post('/api/restore',data={'file':(io.BytesIO(b'not a database'),'bad.sqlite3'),'confirmation':'恢复备份'},headers={'X-CSRF-Token':admin.csrf},content_type='multipart/form-data')
    assert response.status_code==400
    assert admin.call('/bootstrap').json['user']['role']=='leader'


def test_account_last_leader_and_disable(admin):
    identifier=admin.call('/bootstrap').json['user']['id']
    assert admin.call('/users/'+identifier+'/change','POST',{'active':False}).status_code==400


def test_real_xls_can_be_read_when_fixture_available():
    path=Path(r'C:\Users\HONOR\Desktop\客户报价清单\暖通询价单.xls')
    if not path.exists(): pytest.skip('本机客户样本不存在')
    parsed=read_workbook(path.read_bytes(),path.name)
    assert parsed and any(s['rows'] for s in parsed)
