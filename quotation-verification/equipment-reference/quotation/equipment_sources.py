"""Read-only workbook extraction; never use cached finished costs as dynamic costs."""
import hashlib
import re
from pathlib import Path

import openpyxl

from .engine import RuleError, number, text_number, variables
from .equipment_catalog import default_catalog, validate_catalog


def powers(value):
    raw = str(value).strip().upper().replace('KW', '').replace('千瓦', '')
    scale = 1000 if raw.endswith('W') else 1
    raw = raw.removesuffix('W').replace('／', '/')
    if not re.fullmatch(r'\d+(?:\.\d+)?(?:/\d+(?:\.\d+)?)*', raw):
        raise RuleError(f'功率格式待确认：{value}')
    return sorted([text_number(number(x) / scale) for x in raw.split('/')], key=number, reverse=True)


def symbolic(sheet, cell, overrides, trail=()):
    if cell in overrides: return str(overrides[cell])
    if cell in trail: raise RuleError(f'成本公式存在循环：{sheet.title}!{cell}')
    value = sheet[cell].value
    if isinstance(value, (int, float)): return text_number(number(value))
    if not isinstance(value, str) or not value.startswith('='):
        raise RuleError(f'成本公式缺少输入：{sheet.title}!{cell}')
    expression = value[1:].replace('$', '')
    expression = re.sub(r'\b[A-Z]{1,3}\d+\b',
                        lambda m: '(' + symbolic(sheet, m[0], overrides, (*trail, cell)) + ')', expression)
    variables(expression)  # Use the same bounded arithmetic grammar as the runtime.
    return expression


def import_catalog(base_path, pricebook_path, swf_path):
    base_path, pricebook_path, swf_path = map(Path, (base_path, pricebook_path, swf_path))
    catalog = default_catalog()
    for p in (base_path, pricebook_path, swf_path):
        catalog['source_files'].append({'name': p.name, 'sha256': hashlib.sha256(p.read_bytes()).hexdigest()})
    catalog['source'] += '；' + '；'.join(p.name for p in (base_path, pricebook_path, swf_path))
    base = openpyxl.load_workbook(base_path, data_only=False)
    values = openpyxl.load_workbook(base_path, data_only=True)
    latest = openpyxl.load_workbook(pricebook_path, data_only=False)
    swf = openpyxl.load_workbook(swf_path, data_only=True)
    try:
        for label, s in [('base', values['风机如何选型']), ('swf', swf['SWF混流'])]:
            for row in range(2, s.max_row + 1):
                model = str(s.cell(row, 1).value or '').strip()
                if not re.match(r'^(HTFC|HTF\(A\)|SWF)-', model): continue
                family = model.rsplit('-', 1)[0] if model.startswith('HTF(A)') else model.split('-')[0]
                try:
                    point = {'id': f'{label}:{row}', 'model': model, 'family': family,
                             'size': text_number(number(s.cell(row, 2).value)),
                             'group': str(s.cell(row, 3).value or ''),
                             'speed': text_number(number(s.cell(row, 5).value)),
                             'airflow': text_number(number(s.cell(row, 6).value)),
                             'pressure': text_number(number(s.cell(row, 7).value)),
                             'powers': powers(s.cell(row, 8).value),
                             'poles': re.findall(r'\d+', str(s.cell(row, 9).value)),
                             'source': f'{base_path.name if label == "base" else swf_path.name}/{s.title}!A{row}:I{row}'}
                except RuleError as error:
                    raise RuleError(f'选型源行需核实：{s.title}第{row}行：{error}')
                point['quarantined'] = (model == 'HTFC-20' and point['speed'] == '750' and
                                        point['airflow'] == '16536' and point['pressure'] == '18447')
                if point['quarantined']: point['quality_note'] = 'Word明确标记风压18447Pa异常，禁止选用'
                catalog['fan_points'].append(point)
        # New pricebook motor records supersede corresponding old inputs.
        s = latest['新电机价格']
        for row in range(1, s.max_row + 1):
            for col in (1, 4, 7, 10, 13, 16, 19):
                model = str(s.cell(row, col).value or '').strip()
                if not re.match(r'^(YE\d|YD)', model): continue
                cost = s.cell(row, col + 2).value
                motor = {'id': f'motor:{row}:{col}', 'model': model,
                         'powers': powers(s.cell(row, col + 1).value),
                         'poles': re.findall(r'\d+', model.rsplit('-', 1)[-1]),
                         'cost': text_number(number(cost)),
                         'source': f'{pricebook_path.name}/{s.title}!{s.cell(row,col+2).coordinate}'}
                catalog['motors'].append(motor)
        # HTFC source provides shell weight and component formulas, separate from the motor.
        s = base['柜式离心风机HTFC']
        for row in range(3, 11):
            size = text_number(number(s[f'A{row}'].value))
            replace = {f'A{row}': 'size', f'G{row}': 'steel_price'}
            catalog['fan_bodies'].append({
                'id': f'body:HTFC:{size}', 'family': 'HTFC', 'size': size,
                'material_expression': symbolic(s, f'B{row}', replace),
                'other_expression': '+'.join('('+symbolic(s, f'{col}{row}', replace)+')' for col in 'CDEF'),
                'steel_price': text_number(number(s[f'G{row}'].value)), 'thickness': None,
                'source': f'{base_path.name}/{s.title}!A{row}:G{row}'})
        # The latest motor sheet contains the latest axial body formulas (not its filled motor price).
        s = latest['新电机价格']
        for row in range(3, 18):
            if not isinstance(s[f'A{row}'].value, (float, int)): continue
            size = text_number(number(s[f'A{row}'].value))
            formula = s[f'E{row}'].value
            match = re.search(r'\*([\d.]+)\*7\.85', formula or '')
            if not match: raise RuleError(f'轴流机壳板厚未识别：E{row}')
            thickness = match[1]
            expr = symbolic(s, f'E{row}', {f'C{row}': 'steel_price'})
            expr = expr.replace(f'*{thickness}*7.85', '*thickness*7.85', 1)
            replace = {f'A{row}': 'size', f'C{row}': 'steel_price'}
            for family in ('HTF(A)', 'SWF'):
                catalog['fan_bodies'].append({
                    'id': f'body:{family}:{size}', 'family': family, 'size': size,
                    'material_expression': expr,
                    'other_expression': '+'.join('('+symbolic(s, f'{col}{row}', replace)+')' for col in 'FGH'),
                    'steel_price': text_number(number(s[f'C{row}'].value)), 'thickness': thickness,
                    'source': f'{pricebook_path.name}/{s.title}!A{row}:I{row}'})
        s = values['风管类']
        for row in range(14, 26):
            grade = '201' if row < 20 else '304'
            # B is the sequence number; D is the physical sheet thickness in mm.
            thickness = s[f'D{row}'].value
            if not isinstance(thickness, (int, float)):
                match = re.search(r'\d+(?:\.\d+)?', str(thickness))
                if not match: raise RuleError(f'材料厚度未识别：风管类!D{row}')
                thickness = match[0]
            catalog['materials'].append({'id': f'material:{grade}:{row}', 'material': grade,
                'thickness': text_number(number(thickness)), 'ton_price': text_number(number(s[f'E{row}'].value)),
                'source': f'{base_path.name}/风管类!E{row}'})
        s = values['外购产品']
        for row in range(2, s.max_row + 1):
            cost, model = s[f'N{row}'].value, s[f'D{row}'].value
            if not model or not isinstance(cost, (int, float)): continue
            record = {'id': f'purchased:{row}', 'name': str(s[f'B{row}'].value),
                      'model': str(model), 'cost': text_number(number(cost)), 'unit': s[f'M{row}'].value or '台',
                      'source': f'{base_path.name}/外购产品!N{row}'}
            for key, col in [('airflow', 'E'), ('pressure', 'F')]:
                if isinstance(s[f'{col}{row}'].value, (int, float)):
                    record[key] = text_number(number(s[f'{col}{row}'].value))
            catalog['purchased'].append(record)
        validate_catalog(catalog)
        return catalog
    finally:
        for w in (base, values, latest, swf): w.close()
