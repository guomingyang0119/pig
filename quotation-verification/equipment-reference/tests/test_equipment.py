import io
import json
from copy import deepcopy
from decimal import Decimal as D
from pathlib import Path

import openpyxl
import pytest

from quotation.engine import calculate, RuleError, validate_resource
from quotation.equipment import valve_cost, area_cost, parse_attributes
from quotation.equipment_catalog import default_catalog
from quotation.equipment_fans import select_fan, FanSelectionError, fan_cost, material_price
from quotation.excel_io import write_quote, build_quote_rows


@pytest.fixture
def ecatalog():
    b = default_catalog(); b['effective_from'] = '2020-01-01'
    b['purchased'] = [{'id':'purchase', 'model':'TEST', 'name':'测试设备', 'cost':'75', 'unit':'台'}]
    return {'id':'equipment', 'family':'equipment', 'kind':'equipment_catalog', 'status':'published', 'version':1, 'body':b}


def erow(kind='purchased', **attrs):
    return {'name':'测试设备', 'spec':'', 'category':'外购设备', 'unit':'台', 'quantity':'2',
            'attributes':{'equipment_kind':kind, 'model':'TEST', **attrs}}


def equote(row=None, **options):
    return {'pricing_mode':'equipment_v1', 'equipment_options':{'payment_option':'no_install_cash', 'tax_included':False, 'include_freight':False, 'region':'贵州', **options}, 'rows':[row or erow()]}


def computed(ecatalog, row=None, **options):
    return calculate(equote(row, **options), [ecatalog], {'formal_template_confirmed':True})


def test_catalog_needs_publication_and_never_falls_back(ecatalog):
    draft=deepcopy(ecatalog); draft['status']='draft'
    assert not calculate(equote(),[draft],{})['formal_ready']
    retired=deepcopy(ecatalog); retired.update(version=2,status='retired')
    assert calculate(equote(),[ecatalog,retired],{})['rows'][0]['status']=='pending'
    assert computed(ecatalog)['rows'][0]['amount']=='200'


def test_exact_pricing_order_manual_freight_and_upward_rounding(ecatalog):
    ecatalog['body']['purchased'][0]['cost']='75.01'
    row=erow(freight_per_unit='2.2'); row['quantity']='1.5'
    r=computed(ecatalog,row, tax_included=True,payment_option='no_install_days_7',include_freight=True)['rows'][0]
    assert r['unit_price']=='117'  # 75.01/.75*1.04*1.1 + 2.2 = 116.615...
    assert r['amount']=='176'      # ceil(117*1.5)
    assert D(r['supply_unit_unrounded']) < D('114.42')
    assert r['detail']['freight_basis']=='manual_per_unit'


@pytest.mark.parametrize('region,factor',[('贵州','1.03'),('四川','1.05'),('广西','1.05'),('云南','1.05'),('陕西','1.06'),('海南','1.1')])
def test_freight_boundary_is_per_row_unrounded_supply(ecatalog,region,factor):
    row=erow();row['quantity']='80'
    r=computed(ecatalog,row,include_freight=True,region=region)['rows'][0]
    assert D(r['unit_price'])==100*D(factor)
    assert r['detail']['freight_basis']=='regional'
    row['quantity']='79.999'
    assert '8000' in computed(ecatalog,row,include_freight=True)['rows'][0]['reason']
    row['attributes']['freight_per_unit']='0'
    assert computed(ecatalog,row,include_freight=True)['rows'][0]['unit_price']=='100'
    q=equote(row,include_freight=True);q['rows']=[erow(),erow()];q['rows'][0]['quantity']='79'
    assert calculate(q,[ecatalog],{})['pending_count']==2


def test_outside_region_and_unknown_payment_never_free(ecatalog):
    assert computed(ecatalog,erow(freight_per_unit='0'),include_freight=True,region='湖南')['rows'][0]['status']=='pending'
    assert computed(ecatalog,payment_option='')['rows'][0]['status']=='pending'
    assert computed(ecatalog,include_installation=True)['rows'][0]['status']=='pending'


def test_repeating_decimal_at_freight_boundary_is_exact(ecatalog):
    ecatalog['body']['purchased'][0]['cost']='1'
    row=erow();row['quantity']='6000'
    r=computed(ecatalog,row,include_freight=True)['rows'][0]
    assert r['supply_amount_unrounded']=='8000'  # (1 / .75) * 6000, not 7999.999...
    assert r['detail']['freight_basis']=='regional'
    ecatalog['body']['purchased'][0]['cost']='50'
    r=computed(ecatalog,erow(),payment_option='no_install_days_15')['rows'][0]
    assert r['unit_price']=='70'  # (50 / .75) * 1.05 is exactly 70


def test_source_minimum_louver_cost_applied_after_attachment_factor(ecatalog):
    r=computed(ecatalog,erow('louver',width='200',height='120',louver_type='single_valve'))['rows'][0]
    assert D(r['unit_cost'])==24
    assert r['unit_price']=='32'
    ecatalog['body']['louvers']['single_valve'][0]='1000'
    assert D(computed(ecatalog,erow('louver',width='200',height='120',louver_type='single_valve'))['rows'][0]['unit_cost'])==D('93.6')


def test_split_and_splice_are_distinct_and_check_valves_exempt(ecatalog):
    c=ecatalog['body']
    for width,height,split,splice in [('1250','1000',False,0),('1251','1000',True,0),('2000','1000',True,0),('2001','1000',True,40),('1200','1001',True,40)]:
        _,d=valve_cost('fire',{'width':width,'height':height},c)
        assert d['split']==split;assert D(d['components']['splice'])==splice
    with pytest.raises(RuleError,match='拼接'): valve_cost('fire',{'width':'2001','height':'1001'},c)
    _,d=valve_cost('fire',{'width':'2001','height':'1001','splice_cost':'66'},c)
    assert D(d['components']['splice'])==66
    _,d=valve_cost('manual',{'width':'1200','height':'1200'},c)
    assert D(d['components']['splice'])==20
    _,d=valve_cost('check',{'width':'2200','height':'1200'},c)
    assert D(d['components']['splice'])==0


def test_remote_cost_before_margin_and_panel_once(ecatalog):
    a={'width':'400','height':'600'};c=ecatalog['body']
    normal,near=valve_cost('plate_smoke',a,c)
    remote,d=valve_cost('plate_smoke',{**a,'control':'remote','remote':'true'},c)
    assert D(d['components']['mechanism'])==99
    assert remote-normal==77  # 39+60 instead of 22, only once
    assert D(d['components']['panel'])==D('.46')*D('.66')*80
    assert D(d['components']['link'])==D(d['components']['shaft'])==0
    row=erow('plate_smoke',**a,control='remote')
    priced=computed(ecatalog,row)['rows'][0]
    assert D(priced['unit_cost'])==remote
    assert D(priced['unit_price'])==(remote/D('.75')).to_integral_value(rounding='ROUND_CEILING')


def test_current_row_blade_count_and_round_price_before_rounding(ecatalog):
    _,d=valve_cost('fire',{'width':'400','height':'450'},ecatalog['body'])
    assert d['blades']=='3'
    assert D(d['components']['link'])==D('1.5')+D('2')*D('.2')*D('7.85')*D('1.2')*4500*D('.035')/1000
    square=erow('fire',width='500',height='500')
    round_row=erow('fire',diameter='500')
    s=computed(ecatalog,square)['rows'][0];r=computed(ecatalog,round_row)['rows'][0]
    assert D(r['supply_unit_unrounded'])==D(s['supply_unit_unrounded'])*D('1.5')


def test_material_prices_change_only_metal_components(ecatalog):
    c=ecatalog['body'];c['materials']=[{'id':'m','material':'304','thickness':'1.2','ton_price':'9000'}]
    base,a=valve_cost('fire',{'width':'500','height':'500'},c)
    steel,b=valve_cost('fire',{'width':'500','height':'500','material':'304'},c)
    for key in ('frame','blade','brace'): assert D(b['components'][key])==D(a['components'][key])*2
    for key in ('labor','mechanism','corners','shaft','consumables'): assert a['components'][key]==b['components'][key]
    assert steel>base
    with pytest.raises(RuleError): material_price(c,'304','.75',4500)


def test_silencer_dimensions_plenum_alias_and_elbow_fixed_taxed_price(ecatalog):
    row=erow('silencer');row['spec']='320×250×1500'
    r=computed(ecatalog,row)['rows'][0]
    assert D(r['unit_cost'])==D('2.604')*90
    row['name']='接管箱';row['attributes']={};row['spec']='320×250×1500'
    assert D(computed(ecatalog,row)['rows'][0]['unit_cost'])==D('1.87')*100
    row=erow('elbow',area_per_unit='2.5',freight_per_unit='7');row['quantity']='2'
    r=computed(ecatalog,row,tax_included=True,payment_option='no_install_days_7',include_freight=True)['rows'][0]
    assert r['unit_price']=='527' and r['amount']=='1054'
    row['unit']='㎡';row['quantity']='2.5'
    r=computed(ecatalog,row,tax_included=True)['rows'][0]
    assert r['unit_price']=='200' and r['amount']=='500'


def fan_catalog():
    c=default_catalog()
    c['fan_bodies']=[{'id':'b5','family':'HTF(A)','size':'5','material_expression':'steel_price/100','other_expression':'10','steel_price':'4500','thickness':'1'}]
    c['motors']=[{'id':'m','model':'YE2-4','powers':['3'],'poles':['4'],'cost':'100'}]
    def point(id,model,size,q,p):return {'id':id,'model':model,'family':'HTF(A)','size':size,'powers':['3'],'poles':['4'],'speed':'1450','airflow':q,'pressure':p,'group':''}
    c['fan_points']=[point('small','HTF(A)-4','4','1000','100'),point('large','HTF(A)-5','5','2000','200')]
    return c


def test_fan_alias_power_hard_match_and_cost_reselection():
    c=fan_catalog();row={'name':'风机'}
    cost,d=select_fan(row,c,{'model':'HTF-I-5','power':'2.2','airflow':'2000','pressure':'200'})
    assert cost==155 and d['selected_fan']['model']=='HTF(A)-5'
    cost,d=select_fan(row,c,{'model':'HTF-I-4','power':'2.2','airflow':'1000','pressure':'100'})
    assert d['substituted'] and d['selected_fan']['model']=='HTF(A)-5'
    cost,d=select_fan(row,c,{'model':'GXF','airflow':'1000','pressure':'100'})
    assert d['selected_fan']['model']=='HTF(A)-5'
    with pytest.raises(FanSelectionError): select_fan(row,c,{'model':'HTF-I-5','power':'4'})
    with pytest.raises(FanSelectionError): select_fan(row,c,{'model':'HTF-I-4'})
    c['fan_points'].append({**c['fan_points'][1],'id':'another','airflow':'2500'})
    with pytest.raises(FanSelectionError,match='多个'): select_fan(row,c,{'model':'HTF-I-5','airflow':'1500','pressure':'100'})
    _,d=select_fan({**row,'candidate_id':'large'},c,{'model':'HTF-I-5','airflow':'1500','pressure':'100'})
    assert d['selected_fan']['id']=='large'
    c['fan_points'][1]['quarantined']=True
    with pytest.raises(FanSelectionError): select_fan({**row,'candidate_id':'large'},c,{'model':'HTF-I-5','airflow':'1500','pressure':'100'})


def test_dual_motor_keeps_pair_group_and_power_override():
    c=fan_catalog();base=c['fan_points'][1]
    c['fan_points']=[{**base,'id':'dual','model':'HTFC-12-1','family':'HTFC','size':'12','powers':['2.2','1.5'],'poles':['4','6'],'group':'1'}]
    c['fan_bodies']=[{**c['fan_bodies'][0],'id':'body12','family':'HTFC','size':'12'}]
    c['motors']=[{'id':'motorpair','model':'YD-6/4','powers':['1.5','2.2'],'poles':['6','4'],'cost':'500'}]
    cost,detail=select_fan({'name':'双速柜式风机'},c,{'model':'HTFC-12-1','power':'2.2/1.5','airflow':'2000','pressure':'200'})
    assert cost==555 and detail['selected_fan']['group']=='1'
    assert detail['selected_fan']['powers']==['2.2','1.5']
    c['motors'][0]['powers']=['2.2']
    with pytest.raises(FanSelectionError,match='电机成本'):select_fan({'name':'双速风机'},c,{'model':'HTFC-12-1','power':'2.2/1.5'})
    c=fan_catalog();c['fan_points'][1]['powers']=['11'];c['motors'][0]['powers']=['11']
    _,detail=select_fan({'name':'风机','remark':'按11KW单价'},c,{'model':'HTF-I-5','power':'4','airflow':'2000','pressure':'200'})
    assert detail['selected_fan']['powers']==['11'] and detail['power_override']=='11'


def test_formal_equipment_export_scope_and_actual_configuration(ecatalog):
    c=fan_catalog();ecatalog['body']=c;c['effective_from']='2020-01-01'
    row=erow('fan',model='GXF',power='2.2',airflow='1000',pressure='100');row['spec']='GXF 2.2kW'
    q=equote(row,tax_included=True)
    result=calculate(q,[ecatalog],{'formal_template_confirmed':True})
    wb=openpyxl.load_workbook(io.BytesIO(write_quote(q,result,'TEST',True)))
    ws=wb.active
    assert ws['B7'].value=='含税；不含运费；不含安装'
    assert ws['D10'].value.startswith('HTF(A)-5；功率3kW')
    assert ws['H10'].number_format=='0' and ws['I10'].number_format=='0'
    assert '最小合格规格' in ws['J10'].value
    wb.close()


def test_mixed_excluded_group_does_not_require_its_catalog(ecatalog):
    row={'name':'说明','category':'风管','excluded':True,'exclude_reason':'说明行'}
    q=equote();q['pricing_mode']='mixed_v1';q['rows'].append(row)
    result=calculate(q,[ecatalog],{'formal_template_confirmed':True})
    assert result['formal_ready'] and result['rows'][1]['status']=='excluded'


def test_bad_catalog_is_validation_error_not_server_failure(ecatalog):
    for key,value in [('materials',[None]),('rates',[]),('fan_points',[{'id':'x'}])]:
        bad=deepcopy(ecatalog['body']);bad[key]=value
        with pytest.raises(RuleError):validate_resource('equipment_catalog',bad)


@pytest.mark.parametrize('cost,factor',[('1000','1.4'),('1000.01','1.6'),('2000','1.6'),('2000.01','2')])
def test_explosion_tiers_on_cost_not_sale(ecatalog,cost,factor):
    c=fan_catalog();c['motors'][0]['cost']=str(D(cost)-55)
    ecatalog['body']=c;c['effective_from']='2020-01-01'
    r=erow('fan',model='HTF-I-5',power='3',airflow='2000',pressure='200',explosion='true')
    value=computed(ecatalog,r)['rows'][0]
    assert D(value['unit_cost'])==D(cost)*D(factor)


def test_imported_seed_materials_motors_and_quarantined_point():
    c=json.loads((Path(__file__).resolve().parents[1]/'assets/equipment_catalog_20260924.json').read_text(encoding='utf-8'))
    validate_resource('equipment_catalog',c)
    assert len(c['fan_points'])==743 and len(c['motors'])==130
    assert material_price(c,'304','1.2','4500')==14444
    assert material_price(c,'201','.5','4500')==8094
    assert sum(bool(p.get('quarantined')) for p in c['fan_points'])==1
    point=next(p for p in c['fan_points'] if p['id']=='base:56')
    cost,d=fan_cost(point,c,{})
    assert D(d['material_cost'])==324
    assert cost>324


def test_import_preserves_rows_explicit_sizes_and_missing_quantity():
    sheet={'name':'清单','rows':[{'number':2,'cells':['消声器','320×250×1500','2','个'],'issues':[]},{'number':3,'cells':['合计','','',''],'issues':[]}]}
    rows=build_quote_rows(sheet,2,{'name':0,'spec':1,'quantity':2,'unit':3},'equipment_v1')
    assert len(rows)==2 and parse_attributes(rows[0])['length']=='1500'
    rows[0]['spec']='400×300×2000'
    assert parse_attributes(rows[0])['length']=='2000'  # inferred values never become stale overrides
    assert rows[0]['category']=='消声设备' and rows[1]['quantity']==''


def test_real_quote_notation_does_not_parse_motor_power_as_speed():
    a=parse_attributes({'name':'斜流式风机 GXF-9','spec':'L=30894m³/h P=583pa N=7.5KW'})
    assert a['airflow']=='30894' and a['pressure']=='583' and a['power']=='7.5'
    assert 'speed' not in a
    a=parse_attributes({'name':'风机 HL3-2A-NO6A','spec':'L=14000m³/h H=650pa N=5.5KW'})
    assert a['airflow']=='14000' and a['pressure']=='650'
    a=parse_attributes({'name':'正压送风口','spec':'350*（1800+250）'})
    assert a['width']=='350' and a['height']=='1800'
    a=parse_attributes({'name':'阀','spec':'Φ500','attributes':{'diameter':'600','width':'500','height':'500'}})
    assert a['width']==a['height']=='600'


def test_mixed_uses_independent_engines_order_and_exports(ecatalog):
    from test_duct_pricing import catalog, quote
    duct=quote();q=equote();q.update(pricing_mode='mixed_v1',duct_options=duct['duct_options'],rows=[erow(),duct['rows'][0],erow()])
    result=calculate(q,[ecatalog,catalog()],{'formal_template_confirmed':True})
    assert result['formal_ready'],result['issues']
    assert [r['unit_price'] for r in result['rows']]==['100','31.4','100']
    assert D(result['total'])==31800
    for formal in (False,True):
        content=write_quote(q,result,'TEST',formal)
        wb=openpyxl.load_workbook(io.BytesIO(content)); cells=[c.value for ws in wb for row in ws for c in row if c.value is not None]
        assert '测试设备' in cells and '镀锌钢板风管' in cells
        assert not any('unit_cost' in str(x) or '成本' in str(x) for x in cells)
        assert sum(1 for x in cells if x=='测试设备')==2
        wb.close()
    q['equipment_options']['tax_included']=True
    assert not calculate(q,[ecatalog,catalog()],{'formal_template_confirmed':True})['formal_ready']


def test_api_catalog_and_quote_lifecycle(admin,ecatalog):
    seed=admin.call('/equipment/catalog-seed'); assert seed.status_code==200
    draft=admin.call('/resources','POST',{'kind':'equipment_catalog','body':ecatalog['body']}).json['id']
    assert admin.call('/resources/'+draft+'/publish','POST',{}).status_code==200
    duplicate=admin.call('/resources','POST',{'kind':'equipment_catalog','body':ecatalog['body']}).json['id']
    assert admin.call('/resources/'+duplicate+'/publish','POST',{}).status_code==409
    id=admin.call('/quotes','POST',{'pricing_mode':'equipment_v1'}).json['id']
    q=admin.call('/quotes/'+id).json
    r=admin.call('/quotes/'+id,'PUT',{**q['body'],**equote(), 'revision':q['revision']})
    assert r.status_code==200,r.json
    r=admin.call('/quotes/'+id+'/calculate','POST',{'revision':r.json['revision']})
    assert r.status_code==200,r.json
    assert r.json['result']['rows'][0]['unit_price']=='100'
    r=admin.call('/quotes/'+id+'/export','POST',{'revision':r.json['revision'],'kind':'draft'})
    assert r.status_code==200,r.json
    assert admin.call(r.json['download'].removeprefix('/api')).status_code==200
